"""Replace a master's PDF data while retaining independently owned planner details."""

from __future__ import annotations

import copy
import hashlib
import os
import re
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import tostring

import openpyxl
from openpyxl.utils import column_index_from_string

from shogun.services.sectioned_workbook_updater import (
    DEFAULT_METADATA_HEADER,
    DEFAULT_METADATA_SHEET,
    STYLE_PROPERTIES,
    WorkbookLayoutContract,
    _plain,
    preflight_workbook,
)
from shogun.services.structured_transformations import _profile_regex_budget
from shogun.services.workbook_master_maintenance import (
    HISTORY_HEADER,
    HISTORY_SHEET,
    MISSING_COMMENT,
    STATE_HEADER,
    STATE_SHEET,
    _desired_rows,
    _fresh_candidate,
    _json,
    _read_state,
    _write_owned_sheet,
)

PLANNER_SHEET = "Planner details"
PLANNER_HEADER = (
    "Article", "Row identity", "Field", "Value", "Value override", "Status",
    "Schema", "Profile hash", "Plan sheet",
)
_LEGACY_COMMENT = "[Shogun] Source record is absent from the current PDFs; retained for planner review."


def _copy_cell(source: Any, target: Any, *, value: bool = True) -> None:
    if value:
        target.value = source.value
        if isinstance(target.value, str):
            target.data_type = "s"
        for name in STYLE_PROPERTIES:
            setattr(target, name, copy.copy(getattr(source, name)))
    target.comment = copy.copy(source.comment)
    target._hyperlink = copy.copy(source.hyperlink)
    if target.hyperlink:
        target.hyperlink.ref = target.coordinate


def _normal(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _normal(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_normal(item) for item in value]
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return _plain(value)


def _sheet_digest(sheet: Any) -> str:
    """Compare saved content and logical styles, never workbook-local style IDs."""
    styles = {}
    cells = []
    for row in sheet:
        for cell in row:
            style_key = tuple(cell._style) if cell.has_style else ()
            if style_key not in styles:
                style = []
                for name in STYLE_PROPERTIES:
                    item = copy.copy(getattr(cell, name))
                    style.append(tostring(item.to_tree(), encoding="unicode") if hasattr(item, "to_tree") else item)
                styles[style_key] = style
            link = cell.hyperlink
            cells.append([
                cell.coordinate, cell.value, cell.data_type if cell.value is not None else None, styles[style_key],
                [cell.comment.text, cell.comment.author] if cell.comment else None,
                [link.target, link.location, link.tooltip, link.display, link.ref] if link else None,
            ])
    payload = {
        "cells": cells, "state": sheet.sheet_state, "freeze": sheet.freeze_panes,
        "columns": {
            key: {**dict(item), "min": item.min or column_index_from_string(key),
                  "max": item.max or item.min or column_index_from_string(key)}
            for key, item in sheet.column_dimensions.items()
        },
        "rows": {str(key): dict(item) for key, item in sheet.row_dimensions.items()},
    }
    return hashlib.sha256(_json(_normal(payload)).encode()).hexdigest()


def _read_planner_details(book: Any, updater: Any, title: str) -> dict:
    if PLANNER_SHEET not in book:
        return {}
    sheet = book[PLANNER_SHEET]
    if tuple(cell.value for cell in sheet[1]) != PLANNER_HEADER:
        raise ValueError("Planner details already exists but is not a recognized Shogun archive.")
    entries = {}
    for row in sheet.iter_rows(min_row=2):
        article, identity, field, _, override, _, schema, profile_hash, plan_sheet = [cell.value for cell in row]
        if (type(schema) is not int or schema != 1 or profile_hash != updater.profile_hash or plan_sheet != title
                or not isinstance(article, str) or not isinstance(identity, str)
                or not re.fullmatch(r"[a-f0-9]{64}", identity)
                or not isinstance(field, str)
                or not re.fullmatch(r"(?:column:[1-9]\d{0,4}|month:\d{4}/(?:0[1-9]|1[0-2]))", field)
                or field.startswith("column:") and int(field[7:]) > 16384
                or type(override) is not bool or (identity, field) in entries):
            raise ValueError("Invalid planner-detail ownership; the master was not replaced.")
        entries[identity, field] = {
            "section": article, "identity": identity, "field": field,
            "cell": copy.copy(row[3]), "override": override, "status": row[5].value,
        }
    return entries


def _field(contract: Any, column: int) -> str:
    month = contract.col_to_month.get(column)
    return f"month:{month}" if month else f"column:{column}"


def _collect_details(updater: Any, sheet: Any, contract: Any, state: dict, details: dict) -> None:
    owners = {entry["row"]: (identity, entry) for identity, entry in state.items()}
    for row in range(contract.data_start_row, sheet.max_row + 1):
        identity, owner = owners.get(row, (
            hashlib.sha256(_json([sheet.title, "unowned", row]).encode()).hexdigest(), {},
        ))
        for cell in sheet[row]:
            field = _field(contract, cell.column)
            previous = details.get((identity, field))
            copied = copy.copy(cell)
            if copied.comment:
                text = copied.comment.text
                for marker in (MISSING_COMMENT, _LEGACY_COMMENT):
                    text = text.replace("\n" + marker, "").replace(marker, "")
                if text:
                    copied.comment = copy.copy(copied.comment)
                    copied.comment.text = text
                else:
                    copied.comment = None
            written = owner.get("values", {}).get(str(cell.column))
            override = _plain(cell.value) != written
            annotation = copied.comment is not None or copied.hyperlink is not None
            if not override and not annotation:
                if previous and previous["status"] == "Applied":
                    del details[identity, field]  # An active planner value was explicitly cleared.
                continue
            # A note on an unchanged source value owns the annotation, not the PDF value.
            if not override:
                copied.value = None
            details[identity, field] = {
                "section": owner.get(
                    "section", str(sheet.cell(row, contract.section_key_column).value or "(unassigned)"),
                ),
                "identity": identity, "field": field, "cell": copied, "override": override,
                "status": "Not in selected PDFs",
            }


def _protected_columns(updater: Any, contract: Any) -> set[int]:
    columns = {contract.section_key_column, *contract.month_to_col.values()}
    for rule in updater.policy.get("record_rules", []):
        for key in ("reference_column", "fill_blank_quantity_column", "fill_blank_date_column"):
            if rule.get(key):
                columns.add(rule[key])
        specs = [rule[key] for key in ("quantity_spec", "date_spec", "reference_spec") if key in rule]
        columns.update(int(column) for column, spec in rule.get("insert_row_values", {}).items() if spec in specs)
    columns.update(rule["value_column"] for rule in updater.policy.get("section_rules", []) if rule.get("value_column"))
    return columns


def _apply_details(updater: Any, sheet: Any, contract: Any, desired: dict, details: dict) -> tuple[list, list]:
    protected = _protected_columns(updater, contract)
    overlays, events = [], []
    for entry in details.values():
        entry["status"] = "Not in selected PDFs"
        source = desired.get(entry["identity"])
        if source is None:
            continue
        field = entry["field"]
        column = (contract.month_to_col.get(field[6:]) if field.startswith("month:") else int(field[7:]))
        if column is None or column > sheet.max_column:
            entry["status"] = "Outside selected planning months"
            continue
        target = sheet.cell(source["candidate_row"], column)
        can_apply_value = entry["override"] and column not in protected and target.value in (None, "")
        if can_apply_value:
            _copy_cell(entry["cell"], target)
            overlays.append({"row": target.row, "column": column})
            entry["status"] = "Applied"
        else:
            _copy_cell(entry["cell"], target, value=False)
            entry["status"] = "PDF value retained; planner detail archived" if entry["override"] else "Applied"
        if entry["override"] and not can_apply_value:
            events.append({
                "status": "PLANNER_VALUE_ARCHIVED", "confidence": "REVIEW",
                "section_id": entry["section"], "source_identity": entry["identity"], "field": field,
            })
    return overlays, events


def _write_details(book: Any, title: str, updater: Any, details: dict) -> None:
    if PLANNER_SHEET in book:
        del book[PLANNER_SHEET]
    sheet = book.create_sheet(PLANNER_SHEET)
    sheet.append(PLANNER_HEADER)
    ordered = sorted(details.values(), key=lambda item: (item["section"], item["identity"], item["field"]))
    for row, entry in enumerate(ordered, 2):
        sheet.append([
            entry["section"], entry["identity"], entry["field"], None, entry["override"], entry["status"],
            1, updater.profile_hash, title,
        ])
        _copy_cell(entry["cell"], sheet.cell(row, 4))
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for column in ("B", "G", "H", "I"):
        sheet.column_dimensions[column].hidden = True
    for column, width in {"A": 20, "C": 22, "D": 52, "E": 18, "F": 52}.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(name="Arial", size=10, bold=True)
    sheet.freeze_panes = "A2"


def refresh_selected_pdf_master(updater: Any, sections: list, output_path: Path | str, sheet_name: str | None) -> dict:
    output = Path(output_path)
    if updater.template_path.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("This updater supports .xlsx workbooks only.")
    if updater.template_path.resolve() == output.resolve() or (
        output.exists() and os.path.samefile(updater.template_path, output)
    ):
        raise ValueError("Output cannot overwrite the input workbook.")
    book = openpyxl.load_workbook(updater.template_path)
    try:
        title = sheet_name or updater.policy.get("sheet_name") or book.sheetnames[0]
        if title in (PLANNER_SHEET, STATE_SHEET, HISTORY_SHEET):
            raise ValueError("The plan and planner archive must be distinct worksheets.")
        sheet = book[title]
        contract = WorkbookLayoutContract(sheet, updater.policy)
        preflight_workbook(book, sheet)
        metadata_name = updater.policy.get("metadata_sheet_name", DEFAULT_METADATA_SHEET)
        if metadata_name in (PLANNER_SHEET, STATE_SHEET, HISTORY_SHEET, title):
            raise ValueError("The plan, planner archive and ownership sheets must be distinct.")
        state, history = _read_state(updater, book, sheet, contract, metadata_name)
        details = _read_planner_details(book, updater, title)
        _collect_details(updater, sheet, contract, state, details)
        original_rows = sheet.max_row
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".shogun-selected-master-", dir=output.parent) as folder_name:
            folder = Path(folder_name)
            candidate_book, fresh, inserted = _fresh_candidate(updater, sections, sheet, folder)
            try:
                candidate = candidate_book[title]
                current = WorkbookLayoutContract(candidate, updater.policy)
                with _profile_regex_budget(updater.profile_id):
                    desired = _desired_rows(updater, sections, candidate, fresh, inserted, {})
                overlays, detail_events = _apply_details(updater, candidate, current, desired, details)
                metadata = [list(DEFAULT_METADATA_HEADER)]
                ownership = [list(STATE_HEADER)]
                for identity, entry in sorted(desired.items()):
                    values = {column: _plain(value) for column, value in entry["values"].items()}
                    row = entry["candidate_row"]
                    metadata.append([1, title, row, identity, _json(values), updater.profile_hash])
                    ownership.append([
                        1, updater.profile_hash, title, entry["section"], identity, row, entry["kind"], entry["rule"],
                        _json(entry["sort"]), _json(values), candidate.max_row,
                    ])
                for name, rows, visibility in (
                    (metadata_name, metadata, "hidden"), (STATE_SHEET, ownership, "veryHidden"),
                    (HISTORY_SHEET, history or [list(HISTORY_HEADER)], "hidden"),
                ):
                    _write_owned_sheet(candidate_book, name, rows, visibility)
                _write_details(candidate_book, title, updater, details)
                events = [*fresh["events"], *detail_events]
                summary = {
                    **fresh, "mode": "maintain_master", "master_refresh_mode": "selected_pdfs",
                    "baseline_rows": original_rows, "planner_details_count": len(details),
                    "planner_details_sheet": PLANNER_SHEET, "planner_value_cells": overlays,
                    "source_rows_replaced": len(state), "metadata_sheet_name": metadata_name,
                    "status_counts": dict(Counter(event["status"] for event in events)), "events": events,
                    "current_pdf_rows": [
                        {"row": entry["candidate_row"],
                         "values": {col: _plain(value) for col, value in entry["values"].items()}}
                        for entry in desired.values()
                    ],
                    "selected_sheet_digests": {ws.title: _sheet_digest(ws) for ws in candidate_book},
                }
                staged = folder / "selected.xlsx"
                candidate_book.save(staged)
                comparison = compare_selected_pdf_master(updater.template_path, staged, summary)
                if comparison["status"] != "PASS":
                    raise ValueError(f"Selected-PDF master verification failed: {comparison['issues'][:5]}")
                os.replace(staged, output)
                return summary
            finally:
                candidate_book.close()
    finally:
        book.close()


def compare_selected_pdf_master(baseline: Path, generated: Path, summary: dict) -> dict:
    book = openpyxl.load_workbook(generated)
    original = openpyxl.load_workbook(baseline)
    issues = []
    try:
        expected = summary["selected_sheet_digests"]
        if book.sheetnames != list(expected):
            issues.append({"type": "sheet_names"})
        if original.defined_names != book.defined_names:
            issues.append({"type": "defined_names"})
        for name, digest in expected.items():
            if name not in book or _sheet_digest(book[name]) != digest:
                issues.append({"type": "selected_sheet_contents", "sheet": name})
        controlled = {summary["sheet_name"], summary["metadata_sheet_name"], STATE_SHEET, HISTORY_SHEET, PLANNER_SHEET}
        for sheet in original:
            if sheet.title not in controlled and (
                sheet.title not in book or _sheet_digest(sheet) != _sheet_digest(book[sheet.title])
            ):
                issues.append({"type": "unrelated_sheet_changed", "sheet": sheet.title})
        plan = book[summary["sheet_name"]]
        overlays = {(item["row"], item["column"]) for item in summary["planner_value_cells"]}
        for entry in summary["current_pdf_rows"]:
            for column, value in entry["values"].items():
                cell = plan.cell(entry["row"], int(column))
                if (entry["row"], int(column)) not in overlays and _plain(cell.value) != value:
                    issues.append({"type": "pdf_value_changed", "row": entry["row"], "column": column})
        return {
            "status": "FAIL" if issues else "PASS", "issues": issues,
            "baseline_cells_checked": sum(ws.max_row * ws.max_column for ws in original),
            "baseline_rows": summary["baseline_rows"], "generated_rows": plan.max_row,
            "inserted_rows_checked": len(summary["current_pdf_rows"]),
            "planner_details_preserved": summary["planner_details_count"],
            "reference_workbook_status": "NOT_PROVIDED", "reference_row_match_metrics": None,
            "review_required": any(event.get("confidence") == "REVIEW" for event in summary["events"]),
        }
    finally:
        book.close()
        original.close()
