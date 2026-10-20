"""Reconcile current source rows into a planner-maintained generated workbook.

The ordinary population engine remains responsible for parsing and mapping. This
module owns only the persistent row ownership, calendar migration and conservative
cell reconciliation needed between successive source snapshots.
"""

from __future__ import annotations

import copy
import hashlib
import json
import math
import os
import re
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.comments import Comment

from shogun.services.sectioned_workbook_updater import (
    DEFAULT_METADATA_HEADER,
    DEFAULT_METADATA_SHEET,
    STYLE_PROPERTIES,
    SafeWorkbookUpdater,
    WorkbookLayoutContract,
    WorkbookPreservationComparator,
    _plain,
    preflight_workbook,
)
from shogun.services.structured_transformations import _profile_regex_budget, _section_condition_matches

STATE_SHEET = "_shogun_master_state"
HISTORY_SHEET = "_shogun_master_history"
STATE_HEADER = (
    "schema_v1", "profile_hash", "sheet", "section_id", "source_identity", "row",
    "kind", "rule_id", "sort_key", "written_values", "worksheet_rows",
)
HISTORY_HEADER = (
    "schema_v1", "profile_hash", "sheet", "section_id", "source_identity", "source_row", "month", "cell_payload",
)
MISSING_COMMENT = (
    "[Shogun] Previously generated row is not mapped in this run; retained for planner review. "
    "See the current-source audit for scope and date decisions."
)
_LEGACY_MISSING_COMMENT = "[Shogun] Source record is absent from the current PDFs; retained for planner review."


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _row_identity(profile_id: str, section: str, kind: str, rule: str = "") -> str:
    return hashlib.sha256(_json([profile_id, section, kind, rule]).encode("utf-8")).hexdigest()


def _read_state(updater: Any, book: Any, sheet: Any, contract: Any, metadata_name: str) -> tuple[dict, list]:
    metadata = updater._read_metadata(book, sheet, metadata_name, contract)
    if STATE_SHEET not in book:
        if metadata_name in book or HISTORY_SHEET in book or any(
            cell.value is not None for row in sheet.iter_rows(min_row=contract.data_start_row) for cell in row
        ):
            raise ValueError("maintain_master requires a pristine empty seed or a verified generated master.")
        return {}, [list(HISTORY_HEADER)]
    state_sheet = book[STATE_SHEET]
    if tuple(cell.value for cell in state_sheet[1]) != STATE_HEADER or state_sheet.sheet_state != "veryHidden":
        raise ValueError("Unrecognized master ownership sheet; refusing to replace it.")
    result = {}
    seen_rows = set()
    for row in state_sheet.iter_rows(min_row=2, values_only=True):
        if (len(row) != len(STATE_HEADER) or row[0] != 1 or row[1] != updater.profile_hash
                or row[2] != sheet.title or not isinstance(row[3], str)
                or not isinstance(row[4], str) or not re.fullmatch(r"[a-f0-9]{64}", row[4])
                or type(row[5]) is not int or not contract.data_start_row <= row[5] <= sheet.max_row
                or row[5] in seen_rows or row[4] in result
                or row[6] not in {"record", "section_rule", "section"}
                or row[7] is not None and not isinstance(row[7], str)
                or type(row[10]) is not int or row[10] != sheet.max_row
                or str(sheet.cell(row[5], contract.section_key_column).value) != row[3]):
            raise ValueError(
                "Invalid master ownership metadata or changed row structure. "
                "Edit planner cell values and comments; do not insert, delete or reorder master rows."
            )
        try:
            sort_key = json.loads(row[8])
            values = json.loads(row[9])
            if (not isinstance(sort_key, list) or any(
                not isinstance(item, list) or len(item) != 2 or item[0] not in {0, 1, 2}
                or not isinstance(item[1], (str, int, float))
                or isinstance(item[1], float) and not math.isfinite(item[1]) for item in sort_key
            )):
                raise ValueError("Invalid section sort key.")
            if metadata.get(row[4]) != {"row": row[5], "values": values}:
                raise ValueError("Master state disagrees with row provenance.")
            if row[6] == "record":
                rule = next((rule for rule in updater.policy.get("record_rules", [])
                             if (rule.get("id") or rule.get("name") or "record") == row[7]), {})
                reference_column = rule.get("reference_column")
                if (reference_column and str(reference_column) in values
                        and _plain(sheet.cell(row[5], reference_column).value) != values[str(reference_column)]):
                    raise ValueError("Owned order references or row positions changed.")
        except (ValueError, TypeError) as exc:
            raise ValueError("Invalid master ownership values or profile mismatch.") from exc
        seen_rows.add(row[5])
        result[row[4]] = {
            "section": row[3], "row": row[5], "kind": row[6], "rule": row[7] or "",
            "sort": sort_key, "values": values,
        }
    if set(metadata) != set(result):
        raise ValueError("Master ownership must account for every provenance row.")
    if HISTORY_SHEET not in book:
        raise ValueError("Generated master is missing its calendar history.")
    history = [list(row) for row in book[HISTORY_SHEET].iter_rows(values_only=True)]
    if not history or tuple(history[0]) != HISTORY_HEADER:
        raise ValueError("Unrecognized master calendar history.")
    for row in history[1:]:
        if (len(row) != len(HISTORY_HEADER) or row[0] != 1 or row[1] != updater.profile_hash
                or row[2] != sheet.title or type(row[5]) is not int
                or not isinstance(row[6], str) or not re.fullmatch(r"\d{4}/\d{2}", row[6])):
            raise ValueError("Invalid master calendar history.")
        try:
            payload = json.loads(row[7])
            if not isinstance(payload, dict) or "value" not in payload:
                raise ValueError("Missing archived value.")
        except (ValueError, TypeError) as exc:
            raise ValueError("Invalid archived calendar cell.") from exc
    return result, history


def _groups(sheet: Any, contract: Any) -> list[dict]:
    groups = []
    seen = set()
    for row in range(contract.data_start_row, sheet.max_row + 1):
        key = sheet.cell(row, contract.section_key_column).value
        if key is None or str(key).strip() == "":
            continue
        key = str(key).strip()
        if groups and groups[-1]["section"] == key:
            continue
        if key in seen:
            raise ValueError("Master section identifiers must form one contiguous group each.")
        if groups:
            groups[-1]["end"] = row - 1
        groups.append({"section": key, "start": row, "end": sheet.max_row})
        seen.add(key)
    return groups


def _fresh_candidate(updater: Any, sections: list, sheet: Any, folder: Path) -> tuple[Any, dict, dict]:
    candidate_profile = copy.deepcopy(updater.profile)
    candidate_profile["parameters"]["workbook_update"]["mode"] = "populate_template"
    seed = folder / "source-seed.xlsx"
    candidate_path = folder / "source-current.xlsx"
    # Saving a working workbook just to clone it mutates openpyxl's cached column
    # outlines and can change sheetFormatPr on its second save. Keep the master
    # working copy untouched until its single, verified final save.
    seed_book = openpyxl.load_workbook(updater.template_path)
    seed_sheet = seed_book[sheet.title]
    data_start = updater.policy.get("data_start_row", 2)
    styles = [copy.copy(cell._style) for cell in seed_sheet[data_start]]
    row_height = seed_sheet.row_dimensions[data_start].height
    seed_sheet.delete_rows(data_start, seed_sheet.max_row - data_start + 1)
    for row in list(seed_sheet.row_dimensions):
        if row >= data_start:
            del seed_sheet.row_dimensions[row]
    for column, style in enumerate(styles, 1):
        seed_sheet.cell(data_start, column)._style = style
    seed_sheet.row_dimensions[data_start].height = row_height
    for name in (STATE_SHEET, HISTORY_SHEET, updater.policy.get("metadata_sheet_name", DEFAULT_METADATA_SHEET)):
        if name in seed_book:
            del seed_book[name]
    seed_book.save(seed)
    seed_book.close()
    summary = SafeWorkbookUpdater(seed, candidate_profile).execute(sections, candidate_path, sheet.title)
    candidate_book = openpyxl.load_workbook(candidate_path)
    return candidate_book, summary, {item["row"]: item for item in summary["inserted_rows"]}


def _desired_rows(updater: Any, sections: list, candidate: Any, summary: dict, inserted: dict, state: dict) -> dict:
    events_by_row = defaultdict(list)
    for event in summary["events"]:
        if "row" in event:
            events_by_row[event["row"]].append(event)
    section_map = {section.key: section for section in sections}
    desired = {}
    for row, item in inserted.items():
        events = events_by_row[row]
        record = next((event for event in events if "record_id" in event), None)
        stock = next((event for event in events if "source_rule_id" in event), None)
        source = record or stock or next(event for event in events if event["status"] == "CREATED_SECTION")
        section_key = source["section_id"]
        kind = "record" if record else "section_rule" if stock else "section"
        rule = str(source.get("rule_id") or source.get("source_rule_id") or "")
        identity = source["record_id"] if record else _row_identity(updater.profile_id, section_key, kind, rule)
        if identity in desired:
            raise ValueError("Current source has conflicting generated row identities.")
        desired[identity] = {
            "section": section_key, "kind": kind, "rule": rule, "candidate_row": row,
            "sort": [list(value) for value in updater._template_section_sort_key(section_map[section_key])]
            if updater.policy.get("template_section_sort") else [[1, section_key.casefold()]],
            "values": {str(column): candidate.cell(row, int(column)).value for column in item["values"]},
            "event": source,
        }

    # A zero current stock still updates its old generated row even though fresh
    # population intentionally omits new zero-stock rows.
    selected = {item["section"] for item in desired.values()}
    for identity, previous in state.items():
        if previous["kind"] == "section" and identity not in desired and previous["section"] in selected:
            section = section_map[previous["section"]]
            values = {
                str(column): updater._eval_spec(spec, section, None)
                for column, spec in updater.policy.get("template_section_values", {}).items()
            }
            values[str(updater.policy["section_key_column"])] = section.key
            desired[identity] = {**previous, "values": values, "event": None, "candidate_row": None}
        if previous["kind"] != "section_rule" or identity in desired or previous["section"] not in selected:
            continue
        section = section_map[previous["section"]]
        for index, rule in enumerate(updater.policy.get("section_rules", [])):
            rule_id = str(rule.get("id") or rule.get("name") or f"section_rule_{index}")
            if previous["rule"] != rule_id or not _section_condition_matches(section, rule.get("when")):
                continue
            value = updater._eval_spec(rule.get("value_spec"), section, None)
            if value != 0 or not rule.get("zero_value_skip_insert", True):
                continue
            values = {
                str(column): updater._eval_spec(spec, section, None)
                for column, spec in updater.policy.get("template_section_values", {}).items()
            }
            values.update({
                str(column): updater._eval_spec(spec, section, None)
                for column, spec in rule.get("insert_row_values", {}).items()
            })
            if rule.get("value_column"):
                values[str(rule["value_column"])] = value
            values[str(updater.policy["section_key_column"])] = section.key
            desired[identity] = {**previous, "values": values, "event": None, "candidate_row": None}
    # Excel round-trips an explicit empty string as a blank cell. Synthetic
    # section/zero-stock rows must use the same canonical blank representation
    # as rows read back from the ordinary population candidate.
    for entry in desired.values():
        entry["values"] = {column: None if value == "" else value for column, value in entry["values"].items()}
    return desired


def _migrate_calendar(updater: Any, sheet: Any, old: Any, new: Any, state: dict, history: list) -> dict:
    old_months = {month: column for month, column in old.month_to_col.items() if re.fullmatch(r"\d{4}/\d{2}", month)}
    new_months = {month: column for month, column in new.month_to_col.items() if re.fullmatch(r"\d{4}/\d{2}", month)}
    if list(old_months.values()) != list(new_months.values()):
        raise ValueError("Maintained calendars must retain the same dated column layout.")
    if state and old.horizon_start and new.horizon_start and new.horizon_start < old.horizon_start:
        raise ValueError("A maintained master cannot move its calendar backwards.")
    row_owners = {entry["row"]: (identity, entry) for identity, entry in state.items()}
    mapping = {str(column): old_months.get(month) for month, column in new_months.items()}
    if old_months == new_months:
        return mapping
    for row in range(old.data_start_row, sheet.max_row + 1):
        snapshots = {column: copy.copy(sheet.cell(row, column)) for column in old_months.values()}
        for month, column in old_months.items():
            cell = snapshots[column]
            if month not in new_months and (cell.value not in (None, "") or cell.comment or cell.hyperlink):
                identity, owner = row_owners.get(row, ("", {}))
                payload = {"value": _plain(cell.value), "number_format": cell.number_format}
                if cell.comment:
                    payload["comment"] = {"text": cell.comment.text, "author": cell.comment.author}
                if cell.hyperlink:
                    payload["hyperlink"] = {
                        field: getattr(cell.hyperlink, field)
                        for field in ("target", "location", "display", "tooltip")
                    }
                history.append([
                    1, updater.profile_hash, sheet.title,
                    owner.get("section", str(sheet.cell(row, old.section_key_column).value or "")),
                    identity, row, month, _json(payload),
                ])
        for month, column in new_months.items():
            target = sheet.cell(row, column)
            if month in old_months:
                source = snapshots[old_months[month]]
                target.value = source.value
                target._style = copy.copy(source._style)
                target.comment = copy.copy(source.comment)
                target._hyperlink = copy.copy(source.hyperlink)
                if target.hyperlink:
                    target.hyperlink.ref = target.coordinate
            else:
                target.value = None
                target.comment = None
                target._hyperlink = None
        if row in row_owners:
            entry = row_owners[row][1]
            remapped = {}
            for column, value in entry["values"].items():
                old_month = old.col_to_month.get(int(column))
                if old_month in old_months:
                    if old_month in new_months:
                        remapped[str(new_months[old_month])] = value
                else:
                    remapped[column] = value
            entry["values"] = remapped
    return mapping


def _write_owned_sheet(book: Any, name: str, rows: list, state: str) -> dict:
    if name in book:
        del book[name]
    sheet = book.create_sheet(name)
    rows = [[None if value == "" else value for value in row] for row in rows]
    for row in rows:
        sheet.append(row)
    # Store source-derived text literally even when it resembles an Excel formula.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.sheet_state = state
    return {"rows": rows, "state": state}


def maintain_workbook_master(
    updater: Any, sections: list, output_path: Path | str, sheet_name: str | None, options: dict | None,
) -> dict:
    output = Path(output_path)
    if (updater.template_path.resolve() == output.resolve()
            or output.exists() and os.path.samefile(updater.template_path, output)):
        raise ValueError("Output cannot overwrite the input workbook.")
    if updater.template_path.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("This updater supports .xlsx workbooks only.")
    book = openpyxl.load_workbook(updater.template_path)
    try:
        title = sheet_name or updater.policy.get("sheet_name") or book.sheetnames[0]
        if title in {STATE_SHEET, HISTORY_SHEET} or title not in book:
            raise ValueError("Invalid maintained master worksheet.")
        sheet = book[title]
        contract = WorkbookLayoutContract(sheet, updater.policy)
        preflight_workbook(book, sheet)
        original_rows = sheet.max_row
        metadata_name = updater.policy.get("metadata_sheet_name", DEFAULT_METADATA_SHEET)
        if metadata_name in {STATE_SHEET, HISTORY_SHEET, title}:
            raise ValueError("Master state, provenance and planning worksheets must be distinct.")
        state, history = _read_state(updater, book, sheet, contract, metadata_name)
        groups = _groups(sheet, contract)
        original_dimensions = {row: copy.copy(dim) for row, dim in sheet.row_dimensions.items()}
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".shogun-master-", dir=output.parent) as temporary_folder:
            folder = Path(temporary_folder)
            candidate_book, fresh, inserted = _fresh_candidate(updater, sections, sheet, folder)
            try:
                candidate = candidate_book[title]
                current_contract = WorkbookLayoutContract(candidate, updater.policy)
                with _profile_regex_budget(updater.profile_id):
                    desired = _desired_rows(updater, sections, candidate, fresh, inserted, state)
                calendar_mapping = _migrate_calendar(updater, sheet, contract, current_contract, state, history)
                changes, comment_changes = [], []

                def change_cell(row, column, value):
                    cell = sheet.cell(row, column)
                    if _plain(cell.value) != _plain(value):
                        changes.append({
                            "original_row": row, "column": column,
                            "old_value": _plain(cell.value), "new_value": _plain(value),
                        })
                        cell.value = value
                        if isinstance(value, str):
                            cell.data_type = "s"

                for column in current_contract.col_to_month:
                    change_cell(contract.header_row, column, candidate.cell(contract.header_row, column).value)
                events = fresh["events"]
                insertions = defaultdict(list)
                existing_groups = {group["section"]: group for group in groups}
                for event in events:
                    if event["status"] == "CREATED_SECTION" and event["section_id"] in existing_groups:
                        event["status"] = "MATCHED_EXISTING_SECTION"
                        event["original_row"] = existing_groups[event["section_id"]]["start"]
                section_sorts = {}
                for entry in state.values():
                    section_sorts.setdefault(entry["section"], entry["sort"])
                for entry in desired.values():
                    section_sorts.setdefault(entry["section"], entry["sort"])
                candidate_row_targets = {}

                def mark_missing(entry, missing):
                    row, column = entry["row"], contract.section_key_column
                    cell = sheet.cell(row, column)
                    old_comment = cell.comment
                    original = old_comment.text if old_comment else ""
                    text = original
                    for marker in (MISSING_COMMENT, _LEGACY_MISSING_COMMENT):
                        text = text.replace("\n" + marker, "").replace(marker, "")
                    if missing:
                        text = text + ("\n" if text else "") + MISSING_COMMENT
                    if text != original:
                        cell.comment = Comment(text, old_comment.author if old_comment else "Shogun") if text else None
                        comment_changes.append({
                            "original_row": row, "column": column,
                            "text": text or None, "author": cell.comment.author if cell.comment else None,
                        })

                for identity, entry in desired.items():
                    section = entry["section"]
                    if identity not in state:
                        if entry["kind"] == "section" and section in existing_groups:
                            candidate_row_targets[entry["candidate_row"]] = existing_groups[section]["start"]
                            continue
                        if section in existing_groups:
                            anchor = existing_groups[section]["end"]
                        else:
                            following = next((
                                group for group in groups
                                if section_sorts.get(group["section"], [[2, group["section"]]]) > entry["sort"]
                            ), None)
                            anchor = following["start"] - 1 if following else (
                                groups[-1]["end"] if groups else contract.data_start_row - 1
                            )
                        insertions[anchor].append((identity, entry))
                        continue
                    previous = state[identity]
                    row = previous["row"]
                    candidate_row_targets[entry["candidate_row"]] = row
                    written = {}
                    manual = []
                    changed = False
                    for column in set(previous["values"]) | set(entry["values"]):
                        old_value = previous["values"].get(column)
                        new_value = entry["values"].get(column)
                        actual = sheet.cell(row, int(column)).value
                        if _plain(actual) == old_value or column not in previous["values"] and actual in (None, ""):
                            changed |= _plain(actual) != _plain(new_value)
                            change_cell(row, int(column), new_value)
                            written[column] = _plain(new_value)
                        else:
                            written[column] = old_value
                            if _plain(actual) != _plain(new_value):
                                manual.append(int(column))
                    previous["values"] = written
                    mark_missing(previous, False)
                    if entry["event"] is not None and entry["kind"] == "record":
                        entry["event"].update({
                            "status": "PRESERVED_MANUAL_SOURCE_VALUE" if manual else (
                                "UPDATED_SOURCE_RECORD" if changed else "MATCHED_SOURCE_RECORD"
                            ),
                            "confidence": "REVIEW" if manual else "DETERMINISTIC",
                            "original_row": row, "row": row, "preserved_columns": manual,
                        })
                    elif entry["kind"] == "section_rule":
                        status = "PRESERVED_MANUAL_SOURCE_VALUE" if manual else (
                            "UPDATED_SOURCE_ROW" if changed else "MATCHED_SOURCE_ROW"
                        )
                        row_event = entry["event"]
                        if row_event is not None:
                            row_event.update({
                                "status": status, "original_row": row,
                                "confidence": "REVIEW" if manual else "DETERMINISTIC",
                                "preserved_columns": manual,
                            })
                        elif changed or manual:
                            events.append({
                                "status": status, "section_id": section, "source_identity": identity,
                                "source_rule_id": entry["rule"], "original_row": row,
                                "confidence": "REVIEW" if manual else "DETERMINISTIC",
                                "preserved_columns": manual,
                            })
                    elif manual:
                        events.append({
                            "status": "PRESERVED_MANUAL_SOURCE_VALUE", "section_id": section,
                            "source_identity": identity, "original_row": row,
                            "preserved_columns": manual, "confidence": "REVIEW",
                        })
                selected_sections = {entry["section"] for entry in desired.values()}
                source_sections = {section.key for section in sections}
                source_record_ids = {event["record_id"] for event in events if "record_id" in event}
                for identity, entry in state.items():
                    if identity in desired or entry["kind"] == "section" and entry["section"] in selected_sections:
                        continue
                    mark_missing(entry, True)
                    events.append({
                        "status": "MISSING_SOURCE_SECTION_RETAINED" if entry["kind"] == "section" else (
                            "MISSING_SOURCE_RECORD_RETAINED"
                        ),
                        "section_id": entry["section"],
                        "source_identity": identity, "original_row": entry["row"], "confidence": "REVIEW",
                        "current_source_section_present": entry["section"] in source_sections,
                        **({"current_source_record_present": identity in source_record_ids}
                           if entry["kind"] == "record" else {}),
                    })

                def moved(row):
                    return row + sum(len(items) for anchor, items in insertions.items() if anchor < row)

                for entry in state.values():
                    entry["row"] = moved(entry["row"])
                candidate_row_targets = {row: moved(target) for row, target in candidate_row_targets.items()}
                final_insertions = []
                for anchor, items in sorted(insertions.items(), reverse=True):
                    sheet.insert_rows(anchor + 1, len(items))
                    for offset, (identity, entry) in enumerate(items, 1):
                        row = anchor + offset
                        final_row = moved(anchor) + offset
                        candidate_row = entry["candidate_row"]
                        for column in range(1, sheet.max_column + 1):
                            target = sheet.cell(row, column)
                            source = candidate.cell(candidate_row, column)
                            for property_name in STYLE_PROPERTIES:
                                setattr(target, property_name, copy.copy(getattr(source, property_name)))
                            if str(column) in entry["values"]:
                                target.value = entry["values"][str(column)]
                                if isinstance(target.value, str):
                                    target.data_type = "s"
                        state[identity] = {
                            key: entry[key] for key in ("section", "kind", "rule", "sort")
                        } | {"row": final_row, "values": {col: _plain(value) for col, value in entry["values"].items()}}
                        candidate_row_targets[candidate_row] = final_row
                        final_insertions.append({"row": final_row, "values": state[identity]["values"]})
                sheet.row_dimensions.clear()
                for row, dimension in original_dimensions.items():
                    dimension.index = moved(row)
                    sheet.row_dimensions[moved(row)] = dimension
                for item in final_insertions:
                    dimension = copy.copy(candidate.row_dimensions[fresh["data_start_row"] if "data_start_row" in fresh
                                                                    else contract.data_start_row])
                    dimension.index = item["row"]
                    sheet.row_dimensions[item["row"]] = dimension
                for row in sheet:
                    for cell in row:
                        if cell.hyperlink:
                            cell.hyperlink.ref = cell.coordinate

                for event in events:
                    if "original_row" in event:
                        event["row"] = moved(event["original_row"])
                    elif "row" in event and event["row"] in candidate_row_targets:
                        event["row"] = candidate_row_targets[event["row"]]
                metadata_rows = [list(DEFAULT_METADATA_HEADER)]
                state_rows = [list(STATE_HEADER)]
                for identity, entry in sorted(state.items()):
                    metadata_rows.append([
                        1, title, entry["row"], identity, _json(entry["values"]), updater.profile_hash,
                    ])
                    state_rows.append([
                        1, updater.profile_hash, title, entry["section"], identity, entry["row"],
                        entry["kind"], entry["rule"], _json(entry["sort"]), _json(entry["values"]), sheet.max_row,
                    ])
                controlled = {
                    metadata_name: _write_owned_sheet(book, metadata_name, metadata_rows, "hidden"),
                    STATE_SHEET: _write_owned_sheet(book, STATE_SHEET, state_rows, "veryHidden"),
                    HISTORY_SHEET: _write_owned_sheet(book, HISTORY_SHEET, history, "hidden"),
                }
                summary = {
                    **fresh, "mode": "maintain_master", "data_start_row": contract.data_start_row,
                    "baseline_rows": original_rows,
                    "row_mapping": {str(row): moved(row) for row in range(1, original_rows + 1)},
                    "allowed_changes": changes, "allowed_style_changes": [], "allowed_comment_changes": comment_changes,
                    "inserted_rows": final_insertions, "total_rows_inserted": len(final_insertions),
                    "metadata_sheet_name": metadata_name, "controlled_sheets": controlled,
                    "calendar_column_mapping": calendar_mapping,
                    "status_counts": dict(Counter(event["status"] for event in events)),
                    "record_disposition_counts": dict(Counter(
                        event["status"] for event in events if "record_id" in event
                    )),
                }
                staged = folder / "maintained.xlsx"
                book.save(staged)
                comparison = WorkbookPreservationComparator().compare_workbooks(updater.template_path, staged, summary)
                if comparison["status"] != "PASS":
                    raise ValueError(f"Master preservation verification failed: {comparison['issues'][:5]}")
                os.replace(staged, output)
                return summary
            finally:
                candidate_book.close()
    finally:
        book.close()
