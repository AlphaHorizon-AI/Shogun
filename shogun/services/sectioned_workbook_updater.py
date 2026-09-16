"""Profile-driven workbook updating and conservative baseline preservation.

Implements generic in-place row matching, blank-only field fills, section/inventory
reconciliation, reverse-order row insertions, hidden provenance sheets for idempotent
reruns, and cell-by-cell preservation comparison.
"""

from __future__ import annotations

import copy
import datetime
import hashlib
import json
import math
import os
import re
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import coordinate_from_string

from shogun.services.structured_transformations import (
    _canonical_header,
    _compile_pattern,
    _order_sections,
    _planning_column_for_month,
    _planning_header_month,
    _profile_regex_budget,
    _record_matches,
    _resolve_value_spec,
    _section_condition_matches,
)
from shogun.services.transformation_profile_registry import profile_content_hash

DEFAULT_METADATA_SHEET = "_shogun_workflow_provenance"
DEFAULT_METADATA_HEADER = ("schema_v1", "sheet", "row", "source_identity", "written_values", "profile_hash")
MAX_EXCEL_COLUMNS = 16384
MAX_EXCEL_ROWS = 1048576
STYLE_PROPERTIES = (
    "font", "fill", "border", "alignment", "number_format", "protection", "quotePrefix", "pivotButton",
)


def _same_style(original: Any, actual: Any, allowed_change: dict[str, Any] | None = None) -> bool:
    # Saving can deduplicate identical style-table entries. Compare their values,
    # not workbook-local style indexes, while still checking every style property.
    changed_property = allowed_change.get("property") if allowed_change else None
    if allowed_change and changed_property not in STYLE_PROPERTIES:
        return False
    return all(
        copy.copy(getattr(actual, name))
        == (allowed_change["new_value"] if name == changed_property else copy.copy(getattr(original, name)))
        for name in STYLE_PROPERTIES
    )

ALLOWED_WORKBOOK_UPDATE_KEYS = {
    "mode",
    "template_section_values",
    "template_section_sort",
    "sheet_name",
    "header_row",
    "data_start_row",
    "section_key_column",
    "planning_start_column",
    "backlog_headers",
    "future_header_patterns",
    "require_planning_months",
    "metadata_sheet_name",
    "expected_headers",
    "source_identity_fields",
    "section_rules",
    "record_rules",
    "section_matched_status",
    "section_not_found_status",
    "ambiguous_section_status",
    "excluded_section_status",
}

ALLOWED_SECTION_RULE_KEYS = {
    "name",
    "id",
    "description",
    "when",
    "value_spec",
    "value_column",
    "match_column",
    "match_pattern",
    "mixed_match_pattern",
    "zero_value_skip_insert",
    "insert_row_values",
    "update_status",
    "insert_status",
    "policy",
    "update_policy",
}

ALLOWED_RECORD_RULE_KEYS = {
    "id",
    "name",
    "category",
    "match",
    "when",
    "reference_spec",
    "reference_column",
    "match_strip_leading_zeros",
    "quantity_spec",
    "alternative_quantity_spec",
    "fill_blank_quantity_column",
    "date_spec",
    "fill_blank_date_column",
    "date_number_format",
    "planning_month_quantity",
    "legacy_demand_check",
    "insert_row_values",
    "deduplicate_by",
    "insert_status",
    "matched_status",
    "duplicate_status",
    "require_date_in_planning_horizon",
    "require_planning_horizon",
}

ALLOWED_LEGACY_DEMAND_CHECK_KEYS = {
    "blank_column",
    "date_column",
    "check_planning_months",
}

ALLOWED_VALUE_SPEC_KEYS = {
    "literal",
    "section_key",
    "field",
    "group",
    "join",
    "case",
    "coalesce",
    "default",
    "when",
    "value",
    "transforms",
    "value_type",
}


def _plain(value: Any) -> Any:
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat()
    return value.isoformat() if isinstance(value, datetime.date) else value


def _identifier(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _validate_column_index(col: Any, name: str) -> int:
    if isinstance(col, bool) or not isinstance(col, int) or not (1 <= col <= MAX_EXCEL_COLUMNS):
        raise ValueError(f"{name} must be an integer between 1 and {MAX_EXCEL_COLUMNS}.")
    return col


def _validate_row_index(row: Any, name: str) -> int:
    if isinstance(row, bool) or not isinstance(row, int) or not (1 <= row <= MAX_EXCEL_ROWS):
        raise ValueError(f"{name} must be an integer between 1 and {MAX_EXCEL_ROWS}.")
    return row


def _validate_condition_spec(condition: Any) -> None:
    if condition is None:
        return
    if not isinstance(condition, dict):
        raise ValueError("Section condition must be an object.")
    if not condition:
        return
    logical = set(condition) & {"all", "any", "not"}
    if logical:
        if len(condition) != 1:
            raise ValueError("Section condition must have one logical operator.")
        operator = next(iter(logical))
        children = condition[operator] if operator != "not" else [condition[operator]]
        if not isinstance(children, list) or not children:
            raise ValueError("Section condition requires a non-empty condition list.")
        for child in children:
            _validate_condition_spec(child)
        return
    if set(condition) - {"field", "operator", "value", "values", "pattern"}:
        raise ValueError("Unsupported section condition key.")
    operators = {"positive", "equals", "in", "contains", "not_contains", "matches", "not_matches", "truthy"}
    if not isinstance(condition.get("field"), str) or condition.get("operator", "truthy") not in operators:
        raise ValueError("Invalid section condition field or operator.")
    if condition.get("operator") == "in" and not isinstance(condition.get("values"), list):
        raise ValueError("Section 'in' condition requires a values list.")


def _validate_value_spec(spec: Any, context_name: str) -> None:
    if not isinstance(spec, dict):
        raise ValueError(f"{context_name} must be a dictionary.")
    unknown = set(spec.keys()) - ALLOWED_VALUE_SPEC_KEYS
    if unknown:
        raise ValueError(f"Unsupported key(s) in {context_name}: {', '.join(sorted(unknown))}")
    value_types = {
        "string",
        "int",
        "float",
        "number",
        "localized_number",
        "localized_float",
        "strict_localized_number",
        "strict_localized_float",
        "calendar_week_monday",
        "iso_date",
    }
    if "value_type" in spec and spec["value_type"] not in value_types:
        raise ValueError(f"Unsupported value_type in {context_name}.")
    if "transforms" in spec and (
        not isinstance(spec["transforms"], list)
        or any(item not in {"strip", "strip_leading_zero"} for item in spec["transforms"])
    ):
        raise ValueError(f"Unsupported transforms in {context_name}.")
    if "join" in spec:
        join_spec = spec["join"]
        if not isinstance(join_spec, dict):
            raise ValueError(f"{context_name}.join must be an object.")
        allowed_join_keys = {"values", "separator", "require_all"}
        unknown_join = set(join_spec.keys()) - allowed_join_keys
        if unknown_join:
            raise ValueError(f"Unsupported key(s) in {context_name}.join: {', '.join(sorted(unknown_join))}")
        sub_specs = join_spec.get("values")
        if not isinstance(sub_specs, list) or not sub_specs:
            raise ValueError(f"{context_name}.join requires a list of value specs.")
        for sub_idx, sub_spec in enumerate(sub_specs):
            _validate_value_spec(sub_spec, f"{context_name}.join[{sub_idx}]")
    if "case" in spec:
        case_spec = spec["case"]
        if not isinstance(case_spec, list) or not case_spec:
            raise ValueError(f"{context_name}.case must be a non-empty list.")
        for c_idx, candidate in enumerate(case_spec):
            if not isinstance(candidate, dict) or "value" not in candidate:
                raise ValueError(f"{context_name}.case[{c_idx}] must contain 'value'.")
            if set(candidate) - {"when", "value"}:
                raise ValueError(f"Unsupported case key in {context_name}.")
            _validate_condition_spec(candidate.get("when"))
            _validate_value_spec(candidate["value"], f"{context_name}.case[{c_idx}].value")
    if "coalesce" in spec:
        coal_spec = spec["coalesce"]
        if not isinstance(coal_spec, list) or not coal_spec:
            raise ValueError(f"{context_name}.coalesce must be a non-empty list.")
        for c_idx, candidate in enumerate(coal_spec):
            _validate_value_spec(candidate, f"{context_name}.coalesce[{c_idx}]")


def validate_workbook_update_profile(profile: dict[str, Any]) -> None:
    """Bounded, input-independent schema validator for parameters.workbook_update.

    Rejects unsupported keys, invalid types, and out-of-range rows/columns.
    """
    if not isinstance(profile, dict):
        raise ValueError("Profile must be a dictionary.")
    _validate_condition_spec((profile.get("parameters") or {}).get("section_selection"))

    if "workbook_update" in profile:
        policy = profile["workbook_update"]
    elif "parameters" in profile and isinstance(profile["parameters"], dict):
        policy = profile["parameters"].get("workbook_update")
    else:
        policy = profile

    if not isinstance(policy, dict):
        raise ValueError("parameters.workbook_update must be an object.")

    unknown_keys = set(policy.keys()) - ALLOWED_WORKBOOK_UPDATE_KEYS
    if unknown_keys:
        raise ValueError(f"Unsupported key(s) in workbook_update policy: {', '.join(sorted(unknown_keys))}")

    mode = policy.get("mode", "update_existing")
    if not isinstance(mode, str) or mode not in {"update_existing", "populate_template"}:
        raise ValueError("Workbook mode must be 'update_existing' or 'populate_template'.")
    for key in ("backlog_headers", "future_header_patterns"):
        if key in policy:
            values = policy[key]
            if not isinstance(values, list) or any(not isinstance(v, str) or not v.strip() for v in values):
                raise ValueError(f"workbook_update.{key} must be a list of non-empty strings.")
            if key == "future_header_patterns":
                for pattern in values:
                    _compile_pattern(pattern, "future planning header")
    if "template_section_values" in policy:
        values = policy["template_section_values"]
        if not isinstance(values, dict) or not values:
            raise ValueError("template_section_values must be a non-empty column mapping.")
        for column, spec in values.items():
            try:
                column_index = int(column)
            except (ValueError, TypeError):
                raise ValueError("Invalid template section column.")
            _validate_column_index(column_index, "template_section_values")
            _validate_value_spec(spec, "template_section_values")
    if "template_section_sort" in policy:
        specs = policy["template_section_sort"]
        if not isinstance(specs, list) or not 1 <= len(specs) <= 8:
            raise ValueError("template_section_sort requires one to eight value specifications.")
        for spec in specs:
            _validate_value_spec(spec, "template_section_sort")

    header_row = policy.get("header_row", 1)
    header_row = _validate_row_index(header_row, "workbook_update.header_row")

    data_start_row = policy.get("data_start_row", 2)
    data_start_row = _validate_row_index(data_start_row, "workbook_update.data_start_row")
    if data_start_row <= header_row:
        raise ValueError("workbook_update.data_start_row must be greater than header_row.")

    sec_col = policy.get("section_key_column")
    if sec_col is None:
        raise ValueError("workbook_update requires 'section_key_column'.")
    _validate_column_index(sec_col, "workbook_update.section_key_column")

    plan_col = policy.get("planning_start_column")
    if plan_col is not None:
        _validate_column_index(plan_col, "workbook_update.planning_start_column")

    if "require_planning_months" in policy and not isinstance(policy["require_planning_months"], bool):
        raise ValueError("workbook_update.require_planning_months must be a boolean.")

    expected_headers = policy.get("expected_headers")
    if expected_headers is not None:
        if not isinstance(expected_headers, dict):
            raise ValueError("workbook_update.expected_headers must be an object.")
        for k, v in expected_headers.items():
            try:
                col_int = int(k)
            except (ValueError, TypeError):
                raise ValueError(f"Invalid column index in expected_headers: {k!r}")
            _validate_column_index(col_int, f"expected_headers column {k}")
            if not isinstance(v, (str, list, tuple)):
                raise ValueError(f"Expected headers for column {k} must be string or list of strings.")
            if isinstance(v, (list, tuple)) and not all(isinstance(x, str) for x in v):
                raise ValueError(f"Expected headers list for column {k} must contain only strings.")

    source_ident = policy.get("source_identity_fields")
    if source_ident is not None:
        if not isinstance(source_ident, list) or not all(isinstance(spec, dict) for spec in source_ident):
            raise ValueError("workbook_update.source_identity_fields must be a list of value spec objects.")
        for idx, spec in enumerate(source_ident):
            _validate_value_spec(spec, f"source_identity_fields[{idx}]")

    sec_rules = policy.get("section_rules")
    if sec_rules is not None:
        if not isinstance(sec_rules, list):
            raise ValueError("workbook_update.section_rules must be a list of objects.")
        for idx, r in enumerate(sec_rules):
            if not isinstance(r, dict):
                raise ValueError(f"section_rules[{idx}] must be an object.")
            unknown_sec = set(r.keys()) - ALLOWED_SECTION_RULE_KEYS
            if unknown_sec:
                raise ValueError(f"Unsupported key(s) in section_rules[{idx}]: {', '.join(sorted(unknown_sec))}")
            _validate_condition_spec(r.get("when"))
            if "zero_value_skip_insert" in r and not isinstance(r["zero_value_skip_insert"], bool):
                raise ValueError("zero_value_skip_insert must be a boolean.")
            if "policy" in r:
                pol = r["policy"]
                if pol not in {"fill_blank", "replace"}:
                    raise ValueError(
                        f"Unsupported policy in section_rules[{idx}]: {pol!r}. Must be 'fill_blank' or 'replace'."
                    )
            if "update_policy" in r:
                pol = r["update_policy"]
                if pol not in {"fill_blank", "replace"}:
                    raise ValueError(
                        f"Unsupported update_policy in section_rules[{idx}]: {pol!r}. "
                        "Must be 'fill_blank' or 'replace'."
                    )
            if "match_column" in r:
                _validate_column_index(r["match_column"], f"section_rules[{idx}].match_column")
            if "value_column" in r:
                _validate_column_index(r["value_column"], f"section_rules[{idx}].value_column")
            if "value_spec" in r:
                _validate_value_spec(r["value_spec"], f"section_rules[{idx}].value_spec")
            if "insert_row_values" in r:
                ins = r["insert_row_values"]
                if not isinstance(ins, dict):
                    raise ValueError(f"section_rules[{idx}].insert_row_values must be an object.")
                for col_k, col_spec in ins.items():
                    try:
                        col_int = int(col_k)
                    except (ValueError, TypeError):
                        raise ValueError(f"Invalid column in section_rules[{idx}].insert_row_values: {col_k!r}")
                    _validate_column_index(col_int, f"section_rules[{idx}].insert_row_values[{col_k}]")
                    _validate_value_spec(col_spec, f"section_rules[{idx}].insert_row_values[{col_k}]")

    rec_rules = policy.get("record_rules")
    if rec_rules is not None:
        if not isinstance(rec_rules, list):
            raise ValueError("workbook_update.record_rules must be a list of objects.")
        for idx, r in enumerate(rec_rules):
            if not isinstance(r, dict):
                raise ValueError(f"record_rules[{idx}] must be an object.")
            unknown_rec = set(r.keys()) - ALLOWED_RECORD_RULE_KEYS
            if unknown_rec:
                raise ValueError(f"Unsupported key(s) in record_rules[{idx}]: {', '.join(sorted(unknown_rec))}")
            _validate_condition_spec(r.get("when"))
            _record_matches({}, r.get("match"))
            if "match_strip_leading_zeros" in r and not isinstance(r["match_strip_leading_zeros"], bool):
                raise ValueError(f"record_rules[{idx}].match_strip_leading_zeros must be a boolean.")
            if "planning_month_quantity" in r and not isinstance(r["planning_month_quantity"], bool):
                raise ValueError(f"record_rules[{idx}].planning_month_quantity must be a boolean.")
            if "require_date_in_planning_horizon" in r and not isinstance(r["require_date_in_planning_horizon"], bool):
                raise ValueError(f"record_rules[{idx}].require_date_in_planning_horizon must be a boolean.")
            if "require_planning_horizon" in r and not isinstance(r["require_planning_horizon"], bool):
                raise ValueError(f"record_rules[{idx}].require_planning_horizon must be a boolean.")
            if "date_number_format" in r and not isinstance(r["date_number_format"], str):
                raise ValueError(f"record_rules[{idx}].date_number_format must be a string.")
            for col_name in ("reference_column", "fill_blank_quantity_column", "fill_blank_date_column"):
                if col_name in r:
                    _validate_column_index(r[col_name], f"record_rules[{idx}].{col_name}")
            for spec_name in ("reference_spec", "quantity_spec", "alternative_quantity_spec", "date_spec"):
                if spec_name in r:
                    _validate_value_spec(r[spec_name], f"record_rules[{idx}].{spec_name}")
            if "legacy_demand_check" in r:
                chk = r["legacy_demand_check"]
                if not isinstance(chk, dict):
                    raise ValueError(f"record_rules[{idx}].legacy_demand_check must be an object.")
                unknown_chk = set(chk.keys()) - ALLOWED_LEGACY_DEMAND_CHECK_KEYS
                if unknown_chk:
                    raise ValueError(f"Unsupported key(s) in legacy_demand_check: {', '.join(sorted(unknown_chk))}")
                if "check_planning_months" in chk and not isinstance(chk["check_planning_months"], bool):
                    raise ValueError("check_planning_months must be a boolean.")
                if "blank_column" in chk:
                    _validate_column_index(chk["blank_column"], f"record_rules[{idx}].legacy_demand_check.blank_column")
                if "date_column" in chk:
                    _validate_column_index(chk["date_column"], f"record_rules[{idx}].legacy_demand_check.date_column")
            if "insert_row_values" in r:
                ins = r["insert_row_values"]
                if not isinstance(ins, dict):
                    raise ValueError(f"record_rules[{idx}].insert_row_values must be an object.")
                for col_k, col_spec in ins.items():
                    try:
                        col_int = int(col_k)
                    except (ValueError, TypeError):
                        raise ValueError(f"Invalid column in record_rules[{idx}].insert_row_values: {col_k!r}")
                    _validate_column_index(col_int, f"record_rules[{idx}].insert_row_values[{col_k}]")
                    _validate_value_spec(col_spec, f"record_rules[{idx}].insert_row_values[{col_k}]")


class WorkbookLayoutContract:
    """Validate workbook structure and discover dynamic planning month columns."""

    def __init__(self, ws: Any, policy: dict[str, Any]):
        self.header_row = _validate_row_index(policy.get("header_row", 1), "header_row")
        self.data_start_row = _validate_row_index(policy.get("data_start_row", 2), "data_start_row")
        sec_col = policy.get("section_key_column")
        if sec_col is None:
            raise ValueError("Workbook update policy requires 'section_key_column'.")
        self.section_key_column = _validate_column_index(sec_col, "section_key_column")

        plan_col = policy.get("planning_start_column")
        self.planning_start_column = (
            _validate_column_index(plan_col, "planning_start_column") if plan_col is not None else None
        )

        expected_headers = policy.get("expected_headers") or {}
        self.expected_headers = expected_headers
        for col_str, accepted in expected_headers.items():
            col = int(col_str)
            actual_val = str(ws.cell(self.header_row, col).value or "").strip()
            aliases = accepted if isinstance(accepted, (list, tuple)) else [accepted]
            canon_actual = _canonical_header(actual_val)
            if not any(_canonical_header(alias) in canon_actual for alias in aliases):
                raise ValueError(
                    f"Unsupported workbook layout: expected header matching {aliases!r} "
                    f"at column {col}, found {actual_val!r}."
                )

        self.month_to_col: dict[str, int] = {}
        backlog_headers = {str(header).strip().casefold() for header in policy.get("backlog_headers", [])}
        future_patterns = [
            _compile_pattern(pattern, "future planning header")
            for pattern in policy.get("future_header_patterns", [])
        ]
        if self.planning_start_column is not None:
            for col in range(self.planning_start_column, ws.max_column + 1):
                val = ws.cell(self.header_row, col).value
                if val is None:
                    continue
                if str(val).strip().casefold() in backlog_headers:
                    key = "backlog"
                elif isinstance(val, (datetime.date, datetime.datetime)):
                    key = val.strftime("%Y/%m")
                else:
                    text = str(val).strip()
                    month_str = _planning_header_month(text)
                    if month_str:
                        key = month_str
                    else:
                        match = re.fullmatch(r"(\d{4})[-/](\d{2})", text)
                        if match:
                            year, month = map(int, match.groups())
                            key = f"{year:04}/{month:02}"
                        else:
                            try:
                                parsed = datetime.datetime.fromisoformat(text)
                                key = f"{parsed.year:04}/{parsed.month:02}"
                            except ValueError as exc:
                                raise ValueError(f"Unsupported planning header at column {col}: {val!r}") from exc
                    if any(pattern.search(text) for pattern in future_patterns):
                        key = f">={key}"
                if key in self.month_to_col:
                    raise ValueError(f"Duplicate planning month header: {key}")
                self.month_to_col[key] = col

        if not self.month_to_col and policy.get("require_planning_months", False):
            raise ValueError("No planning month headers found in workbook.")
        self.col_to_month = {col: key for key, col in self.month_to_col.items()}
        exact_months = [key for key in self.month_to_col if re.fullmatch(r"\d{4}/\d{2}", key)]
        self.horizon_start = min(exact_months) if exact_months else None
        self.horizon_end = max(exact_months) if exact_months else None


def preflight_workbook(book: Any, sheet: Any) -> None:
    """openpyxl does not shift dependent structures when inserting rows.

    Reject richer templates upfront to prevent silent data/formula corruption.
    """
    unsupported = []
    for ws in book:
        if any(cell.data_type == "f" for row in ws for cell in row):
            unsupported.append(f"formulas in {ws.title}")
        if ws.tables or ws.merged_cells or ws.conditional_formatting or ws.data_validations.count:
            unsupported.append(f"tables/merges/conditional formats/validation in {ws.title}")
        if getattr(ws, "_charts", None) or getattr(ws, "_images", None) or getattr(ws, "_pivots", None):
            unsupported.append(f"drawings or pivots in {ws.title}")
    if any(name.attr_text != "#REF!" for name in book.defined_names.values()) or getattr(book, "_external_links", None):
        unsupported.append("defined names or external links")
    if sheet.print_area or sheet.print_title_rows or sheet.print_title_cols or sheet.auto_filter.ref:
        unsupported.append("print ranges or auto-filter")
    if sheet.row_breaks.brk or sheet.col_breaks.brk:
        unsupported.append("manual page breaks")
    if unsupported:
        raise ValueError("Template requires a structure-aware writer: " + "; ".join(unsupported))


def _compute_source_identity(
    sec: Any,
    rec: Any,
    identity_specs: list[Any] | None = None,
    *,
    profile_id: str = "",
    rule_id: str = "",
) -> str:
    raw_dict = getattr(rec, "raw", None) or (rec if isinstance(rec, dict) else {})
    if identity_specs:
        items = []
        for spec in identity_specs:
            val = _resolve_value_spec(spec, sec, raw_dict)
            items.append(val)
        payload = [profile_id, rule_id, *items]
        return hashlib.sha256(json.dumps(payload, ensure_ascii=False).encode()).hexdigest()

    cleaned_items = sorted(
        (str(k), _plain(v))
        for k, v in raw_dict.items()
        if str(k) not in {"source_file", "source_page", "source_line", "source_text_hash", "line_idx"}
    )
    payload = [profile_id, rule_id, str(getattr(sec, "key", "")), cleaned_items]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, default=str).encode()).hexdigest()


def _build_provenance_record(
    sec: Any,
    rec: Any = None,
    *,
    profile_id: str = "",
    rule_id: str = "",
    identity_specs: list[Any] | None = None,
) -> dict[str, Any]:
    res: dict[str, Any] = {
        "section_id": getattr(sec, "key", ""),
        "source_file": getattr(sec, "source_file", ""),
        "source_pages": getattr(sec, "source_pages", []),
        "confidence": "DETERMINISTIC",
    }
    if rec is not None:
        rec_id = getattr(rec, "identity", "") or _compute_source_identity(
            sec, rec, identity_specs, profile_id=profile_id, rule_id=rule_id
        )
        res.update(
            {
                "record_id": rec_id,
                "category": getattr(rec, "category", ""),
                "reference": getattr(rec, "reference", ""),
                "source_file": getattr(rec, "source_file", ""),
                "source_page": getattr(rec, "source_page", 0),
                "source_line": getattr(rec, "source_line", 0),
                "source_text_hash": getattr(rec, "source_text_hash", ""),
                "chosen_date": _plain(getattr(rec, "date", None)),
                "chosen_quantity": getattr(rec, "quantity", None),
                "raw": dict(getattr(rec, "raw", {})),
            }
        )
    return res


class SafeWorkbookUpdater:
    """Executes profile-driven updates against an existing workbook with 100% preservation."""

    def __init__(self, template_path: Path | str, profile: dict[str, Any]):
        self.template_path = Path(template_path)
        self.profile = profile
        self.profile_hash = profile_content_hash(profile)
        self.profile_id = str(profile.get("id") or "workbook_updater")
        self.parameters = profile.get("parameters") or {}
        self.policy = self.parameters.get("workbook_update") or {}
        validate_workbook_update_profile(profile)

    def execute(
        self,
        sections: list[Any],
        output_path: Path | str,
        sheet_name: str | None = None,
        options: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        output_path = Path(output_path)
        if self.template_path.resolve() == output_path.resolve() or (
            output_path.exists() and os.path.samefile(self.template_path, output_path)
        ):
            raise ValueError("Output cannot overwrite the input workbook.")
        if self.template_path.suffix.lower() != ".xlsx" or output_path.suffix.lower() != ".xlsx":
            raise ValueError("This updater supports .xlsx workbooks only.")

        options = dict(options or {})
        book = openpyxl.load_workbook(self.template_path, data_only=False)
        temporary: Path | None = None
        try:
            target_sheet_name = sheet_name or self.policy.get("sheet_name") or book.sheetnames[0]
            if target_sheet_name not in book.sheetnames:
                raise ValueError(f"Workbook does not contain target sheet '{target_sheet_name}'.")
            ws = book[target_sheet_name]

            metadata_sheet_name = self.policy.get("metadata_sheet_name", DEFAULT_METADATA_SHEET)
            if ws.title == metadata_sheet_name:
                raise ValueError("The selected sheet is provenance metadata, not a planning sheet.")

            contract = WorkbookLayoutContract(ws, self.policy)
            preflight_workbook(book, ws)

            original_rows = ws.max_row
            original_columns = ws.max_column
            populate_template = self.policy.get("mode", "update_existing") == "populate_template"
            if populate_template:
                if any(
                    cell.value is not None
                    for row in ws.iter_rows(min_row=contract.data_start_row)
                    for cell in row
                ):
                    raise ValueError(
                        "Template population requires an empty data area below the headers. "
                        "Select the empty template, or use an update-existing profile for a populated workbook."
                    )
                if metadata_sheet_name in book:
                    raise ValueError("Template population requires a fresh template without generated provenance.")
            original_dimensions = {row: copy.copy(dim) for row, dim in ws.row_dimensions.items()}
            freeze = ws.freeze_panes
            metadata = self._read_metadata(book, ws, metadata_sheet_name, contract)
            metadata_rows = {entry["row"] for entry in metadata.values()}

            sec_col = contract.section_key_column
            blocks: dict[str, list[list[int]]] = defaultdict(list)
            previous = None
            for row in range(contract.data_start_row, original_rows + 1):
                sec_id = _identifier(ws.cell(row, sec_col).value)
                if sec_id:
                    if sec_id != previous:
                        blocks[sec_id].append([])
                    blocks[sec_id][-1].append(row)
                previous = sec_id

            events: list[dict[str, Any]] = []
            insertions: dict[int, list[dict[str, Any]]] = defaultdict(list)
            changes: list[dict[str, Any]] = []
            allowed_style_changes: dict[tuple[int, int], dict[str, Any]] = {}

            selection_cond = self.parameters.get("section_selection")
            relevant = []
            for s in sections:
                if selection_cond:
                    if _section_condition_matches(s, selection_cond):
                        relevant.append(s)
                else:
                    relevant.append(s)

            relevant_ids = {id(s) for s in relevant}

            def section_key_fn(section):
                return section.key

            counts = Counter(section_key_fn(s) for s in relevant)

            ident_specs = self.policy.get("source_identity_fields")

            def event(sec, status, rec=None, **extra):
                rule_id = getattr(rec, "rule_id", "") if rec else ""
                item = {
                    **_build_provenance_record(
                        sec,
                        rec,
                        profile_id=self.profile_id,
                        rule_id=rule_id,
                        identity_specs=ident_specs,
                    ),
                    "status": status,
                    **extra,
                }
                events.append(item)
                return item

            def update_cell(row, col, value, style_prop=None, style_val=None):
                old = ws.cell(row, col).value
                if _plain(old) != _plain(value):
                    changes.append(
                        {
                            "original_row": row,
                            "column": col,
                            "old_value": _plain(old),
                            "new_value": _plain(value),
                        }
                    )
                    ws.cell(row, col).value = value
                    if isinstance(value, str):
                        ws.cell(row, col).data_type = "s"
                if style_prop and style_val:
                    setattr(ws.cell(row, col), style_prop, style_val)
                    allowed_style_changes[(row, col)] = {
                        "original_row": row,
                        "column": col,
                        "property": style_prop,
                        "new_value": style_val,
                    }

            def schedule_insert(sec, anchor, values, status, rec=None, date_format="dd.mm.yyyy"):
                item = event(sec, status, rec)
                # Excel reads an explicitly written empty string back as a blank.
                # New rows already start blank; retain only values that round-trip.
                values = {column: value for column, value in values.items() if value not in (None, "")}
                insertions[anchor].append({"values": values, "event": item, "date_format": date_format})

            # Evaluate updates within independent profile regex budget
            with _profile_regex_budget(self.profile_id):
                if populate_template and self.policy.get("template_section_sort"):
                    sections = sorted(sections, key=self._template_section_sort_key)
                if populate_template and self.parameters.get("section_order"):
                    selected = [section for section in sections if id(section) in relevant_ids]
                    if any(count != 1 for count in counts.values()):
                        raise ValueError("Template section ordering requires unique selected section keys.")
                    ordered = _order_sections(selected, self.parameters, self.profile_id)
                    if Counter(map(id, ordered)) != Counter(map(id, selected)):
                        raise ValueError("Template section ordering must retain every selected section.")
                    # Selection still controls inclusion. Keep excluded sections so
                    # their source records receive an explicit audit disposition.
                    sections = ordered + [section for section in sections if id(section) not in relevant_ids]
                for s in sections:
                    s_key = section_key_fn(s)
                    for exc in getattr(s, "exceptions", []):
                        event(
                            s,
                            exc.get("status", exc.get("type", "SOURCE_RECONCILIATION_WARNING")),
                            detail=exc,
                            confidence="REVIEW",
                        )
                    for page in getattr(s, "source_pages", [])[1:]:
                        event(s, "CONTINUATION_PAGE_ATTACHED", source_page=page)

                    all_recs = getattr(s, "records", getattr(s, "all_records", []))

                    sec_matched_status = self.policy.get("section_matched_status", "MATCHED_SECTION")
                    sec_not_found_status = self.policy.get("section_not_found_status", "SECTION_NOT_FOUND")
                    ambig_sec_status = self.policy.get("ambiguous_section_status", "AMBIGUOUS_SECTION_GROUP")
                    excl_sec_status = self.policy.get("excluded_section_status", "EXCLUDED_SECTION")

                    if id(s) not in relevant_ids:
                        event(s, excl_sec_status, reason="Section outside configured selection rule")
                        for rec in all_recs:
                            event(s, "EXCLUDED_RECORD", rec)
                        continue

                    unresolved_fields = [
                        state for state in getattr(s, "resolution_states", [])
                        if state.get("requires_manual_validation")
                    ]
                    ambiguous_fields = [
                        target for target, outcome in getattr(s, "selector_outcomes", {}).items()
                        if outcome == "ambiguous"
                    ]
                    if unresolved_fields or ambiguous_fields:
                        event(s, "SOURCE_FIELD_REVIEW", confidence="REVIEW",
                              resolution_states=unresolved_fields, ambiguous_fields=ambiguous_fields)

                    groups = blocks.get(s_key, [])
                    if counts[s_key] != 1 or (len(groups) != 1 and not populate_template):
                        status = sec_not_found_status if not groups and not populate_template else ambig_sec_status
                        event(s, status, confidence="REVIEW")
                        for rec in all_recs:
                            event(s, status, rec, confidence="REVIEW")
                        continue

                    if populate_template:
                        rows = []
                        anchor = contract.data_start_row - 1
                        section_values = {
                            int(column): self._eval_spec(spec, s, None)
                            for column, spec in self.policy.get("template_section_values", {}).items()
                        }
                        # Keep each source section visible, including sections whose
                        # records need review or which have no stock or orders.
                        section_values[sec_col] = s_key
                        schedule_insert(s, anchor, section_values, "CREATED_SECTION")
                    else:
                        rows = groups[0]
                        anchor = rows[-1]
                        event(s, sec_matched_status, original_row=rows[0])

                    # ── In-place section rules evaluation ───────────────────────
                    sec_rules = self.policy.get("section_rules") or []
                    for sec_rule in sec_rules:
                        if "when" in sec_rule and not _section_condition_matches(s, sec_rule["when"]):
                            continue

                        val_spec = sec_rule.get("value_spec")
                        sec_val = None
                        if val_spec is not None:
                            raw_val = self._eval_spec(val_spec, s, None)
                            if raw_val is not None and raw_val != "":
                                try:
                                    sec_val = float(raw_val)
                                except (ValueError, TypeError):
                                    sec_val = None
                            if sec_val is None or not math.isfinite(sec_val):
                                event(
                                    s,
                                    "INVALID_SECTION_VALUE",
                                    confidence="REVIEW",
                                    reason=f"Invalid numeric value for section rule '{sec_rule.get('name', '')}'",
                                )
                                continue

                        match_col = sec_rule.get("match_column")
                        val_col = sec_rule.get("value_column")
                        if match_col is None or val_col is None:
                            continue

                        ctx = dict(getattr(s, "fields", {}))
                        ctx["section_key"] = s_key
                        raw_pat = sec_rule.get("match_pattern", "")

                        def substitute_field(match):
                            key = match.group(1)
                            if key not in ctx:
                                raise ValueError(f"Section match pattern references missing field '{key}'.")
                            return re.escape(str(ctx[key]))

                        # Replace named fields while preserving regex quantifiers such as {3}.
                        match_pat = re.sub(r"\{([A-Za-z_]\w*)\}", substitute_field, raw_pat)

                        compiled_match = _compile_pattern(match_pat, f"section_rule_{sec_rule.get('name', 'match')}")
                        matched_rows = [
                            r for r in rows if compiled_match.fullmatch(_identifier(ws.cell(r, match_col).value))
                        ]
                        mixed_pat = sec_rule.get("mixed_match_pattern")
                        if mixed_pat:
                            compiled_mixed = _compile_pattern(
                                mixed_pat, f"section_rule_{sec_rule.get('name', 'mixed')}"
                            )
                            mixed_rows = [
                                r for r in rows if compiled_mixed.search(_identifier(ws.cell(r, match_col).value))
                            ]
                        else:
                            mixed_rows = []

                        upd_status = sec_rule.get("update_status", "UPDATED_SECTION_ROW")
                        ins_status = sec_rule.get("insert_status", "ADDED_SECTION_ROW")
                        update_pol = sec_rule.get("update_policy") or sec_rule.get("policy") or "fill_blank"

                        if len(matched_rows) == 1:
                            r = matched_rows[0]
                            old = ws.cell(r, val_col).value
                            if update_pol == "fill_blank":
                                if old is not None and str(old).strip() != "":
                                    event(
                                        s,
                                        "PRESERVED_EXISTING_SECTION_VALUE",
                                        original_row=r,
                                        preserved_value=_plain(old),
                                        source_value=sec_val,
                                    )
                                else:
                                    update_cell(r, val_col, sec_val)
                                    event(s, upd_status, original_row=r, old_value=old, new_value=sec_val)
                            elif update_pol == "replace":
                                update_cell(r, val_col, sec_val)
                                event(s, upd_status, original_row=r, old_value=old, new_value=sec_val)
                        elif not matched_rows and mixed_rows:
                            event(s, "AMBIGUOUS_SECTION_ROWS", confidence="REVIEW")
                        elif not matched_rows and sec_val == 0 and sec_rule.get("zero_value_skip_insert", True):
                            event(s, "ZERO_VALUE_NO_ROW", reason="No section row exists and current value is zero")
                        elif not matched_rows:
                            insert_vals = {}
                            for c_str, spec in sec_rule.get("insert_row_values", {}).items():
                                insert_vals[int(c_str)] = self._eval_spec(spec, s, None)
                            schedule_insert(s, anchor, insert_vals, ins_status)
                        else:
                            event(s, "AMBIGUOUS_SECTION_ROWS", confidence="REVIEW")

                    # ── Record rules evaluation ─────────────────────────────────
                    rec_rules_list = self.policy.get("record_rules") or []

                    # First pass: match rules and compute identities upfront to identify duplicate source records
                    rec_candidates = []
                    for r in all_recs:
                        raw_record = getattr(r, "raw", {})
                        matching_rules = [
                            rule_cand
                            for rule_cand in rec_rules_list
                            if _record_matches(raw_record, rule_cand.get("match"))
                            and _section_condition_matches(s, rule_cand.get("when"))
                        ]
                        rule_id = (
                            (matching_rules[0].get("id") or matching_rules[0].get("name") or "record")
                            if len(matching_rules) == 1
                            else ""
                        )
                        ident = getattr(r, "identity", "") or _compute_source_identity(
                            s, r, ident_specs, profile_id=self.profile_id, rule_id=rule_id
                        )
                        rec_candidates.append((matching_rules, rule_id, ident))

                    identity_counts = Counter(ident for _, _, ident in rec_candidates if ident)

                    for r, (matching_rules, rule_id, identity) in zip(all_recs, rec_candidates):
                        if len(matching_rules) == 0:
                            event(s, "UNMATCHED_RECORD_RULE", r, confidence="REVIEW")
                            continue
                        elif len(matching_rules) > 1:
                            event(s, "AMBIGUOUS_RECORD_RULE", r, confidence="REVIEW")
                            continue

                        matched_rule = matching_rules[0]
                        r.rule_id = rule_id

                        if getattr(r, "is_invalid", False):
                            event(
                                s,
                                "INVALID_SOURCE_RECORD",
                                r,
                                confidence="REVIEW",
                                reason=getattr(r, "invalid_reason", None),
                            )
                            continue

                        # Check for duplicate source records with identical identity
                        if identity and identity_counts[identity] > 1:
                            dup_status = matched_rule.get("duplicate_status", "DUPLICATE_RECORD_REFERENCE")
                            event(
                                s,
                                dup_status,
                                r,
                                confidence="REVIEW",
                                reason="Duplicate source record with identical identity",
                            )
                            continue

                        # Check reference
                        ref = getattr(r, "reference", "")
                        if "reference_spec" in matched_rule and not ref:
                            ref = str(self._eval_spec(matched_rule["reference_spec"], s, r) or "").strip()

                        has_ref_matching = "reference_column" in matched_rule
                        if has_ref_matching and not ref:
                            event(s, "AMBIGUOUS_RECORD_IDENTITY", r, confidence="REVIEW")
                            continue

                        # Date and quantity
                        rec_date = None
                        if "date_spec" in matched_rule:
                            try:
                                rec_date = self._eval_spec(matched_rule["date_spec"], s, r)
                            except (ValueError, TypeError):
                                pass
                            if not isinstance(rec_date, datetime.date):
                                event(s, "AMBIGUOUS_DATE_MAPPING", r, confidence="REVIEW")
                                continue

                        month_key = (
                            rec_date.strftime("%Y/%m")
                            if isinstance(rec_date, (datetime.date, datetime.datetime))
                            else ""
                        )
                        planning_column = _planning_column_for_month(contract.month_to_col, month_key)

                        # Evaluate quantity and alternative quantity at the writer boundary from raw record
                        qty = None
                        if "quantity_spec" in matched_rule:
                            try:
                                raw_q = self._eval_spec(matched_rule["quantity_spec"], s, r)
                                qty = float(raw_q) if raw_q is not None and raw_q != "" else None
                            except (ValueError, TypeError):
                                qty = None
                        else:
                            qty = getattr(r, "quantity", None)

                        alt_qty = None
                        if "alternative_quantity_spec" in matched_rule:
                            try:
                                raw_alt = self._eval_spec(matched_rule["alternative_quantity_spec"], s, r)
                                alt_qty = float(raw_alt) if raw_alt is not None and raw_alt != "" else None
                            except (ValueError, TypeError):
                                alt_qty = None

                        if "quantity_spec" in matched_rule and (qty is None or not math.isfinite(qty)):
                            event(s, "AMBIGUOUS_QUANTITY_MAPPING", r, confidence="REVIEW")
                            continue
                        if "alternative_quantity_spec" in matched_rule and (
                            alt_qty is None or not math.isfinite(alt_qty)
                        ):
                            event(s, "AMBIGUOUS_QUANTITY_MAPPING", r, confidence="REVIEW")
                            continue
                        if (
                            "quantity_spec" in matched_rule
                            and "alternative_quantity_spec" in matched_rule
                            and qty is not None
                            and alt_qty is not None
                            and qty != alt_qty
                        ) or getattr(r, "is_ambiguous_quantity", False):
                            event(s, "AMBIGUOUS_QUANTITY_MAPPING", r, confidence="REVIEW")
                            continue

                        # If planning_month_quantity is required, ensure date and finite quantity
                        if matched_rule.get("planning_month_quantity"):
                            if rec_date is None:
                                event(s, "AMBIGUOUS_DATE_MAPPING", r, confidence="REVIEW")
                                continue
                            if qty is None or not math.isfinite(qty):
                                event(s, "AMBIGUOUS_QUANTITY_MAPPING", r, confidence="REVIEW")
                                continue

                        # Planning horizon filtering (separate from month-cell writing)
                        require_horizon = (
                            matched_rule.get("require_date_in_planning_horizon")
                            or matched_rule.get("require_planning_horizon")
                            or matched_rule.get("planning_month_quantity")
                        )
                        if require_horizon and not contract.month_to_col:
                            raise ValueError("Record rule requires planning month headers.")
                        if (
                            require_horizon
                            and contract.month_to_col
                            and month_key
                            and planning_column is None
                        ):
                            event(s, "OUTSIDE_PLANNING_HORIZON", r)
                            continue

                        # Build row values for insertion
                        row_values: dict[int, Any] = {}
                        insert_specs = matched_rule.get("insert_row_values") or {}
                        for c_str, spec in insert_specs.items():
                            row_values[int(c_str)] = self._eval_spec(spec, s, r)

                        if matched_rule.get("planning_month_quantity") and planning_column is not None:
                            row_values[planning_column] = qty

                        # Check metadata sheet for rerun idempotency
                        if identity in metadata:
                            entry = metadata[identity]
                            unchanged = all(
                                _plain(ws.cell(entry["row"], int(col)).value) == val
                                for col, val in entry["values"].items()
                            )
                            source_unchanged = all(
                                _plain(row_values.get(int(col))) == val for col, val in entry["values"].items()
                            )
                            status = (
                                "MATCHED_SOURCE_RECORD"
                                if unchanged and source_unchanged
                                else "PRESERVED_SOURCE_CHANGE_REVIEW"
                            )
                            event(
                                s,
                                status,
                                r,
                                original_row=entry["row"],
                                chosen_quantity=qty,
                                confidence="DETERMINISTIC" if unchanged and source_unchanged else "REVIEW",
                            )
                            continue

                        if has_ref_matching:
                            ref_col = int(matched_rule["reference_column"])
                            matches = [
                                row_idx for row_idx in rows if _identifier(ws.cell(row_idx, ref_col).value) == ref
                            ]
                            if not matches and ref.isdigit() and matched_rule.get("match_strip_leading_zeros", False):
                                matches = [
                                    row_idx
                                    for row_idx in rows
                                    if _identifier(ws.cell(row_idx, ref_col).value).isdigit()
                                    and _identifier(ws.cell(row_idx, ref_col).value).lstrip("0") == ref.lstrip("0")
                                ]
                            if len(matches) > 1:
                                event(
                                    s,
                                    matched_rule.get("duplicate_status", "DUPLICATE_RECORD_REFERENCE"),
                                    r,
                                    confidence="REVIEW",
                                )
                                continue
                            if matches:
                                target_r = matches[0]
                                qty_col = matched_rule.get("fill_blank_quantity_column")
                                date_col = matched_rule.get("fill_blank_date_column")
                                old_qty = ws.cell(target_r, qty_col).value if qty_col else None
                                if qty_col and old_qty is None and qty is not None:
                                    update_cell(target_r, qty_col, qty)
                                if date_col and ws.cell(target_r, date_col).value is None and rec_date is not None:
                                    date_fmt = str(matched_rule.get("date_number_format", "dd.mm.yyyy"))
                                    update_cell(
                                        target_r, date_col, rec_date, style_prop="number_format", style_val=date_fmt
                                    )
                                matched_st = matched_rule.get("matched_status", "MATCHED_RECORD")
                                event(
                                    s,
                                    "PRESERVED_MANUAL_QUANTITY"
                                    if old_qty is not None and old_qty != qty
                                    else matched_st,
                                    r,
                                    original_row=target_r,
                                    old_value=old_qty,
                                    new_value=ws.cell(target_r, qty_col).value if qty_col else None,
                                    chosen_quantity=qty,
                                )
                                continue
                        else:
                            legacy_chk = matched_rule.get("legacy_demand_check")
                            if legacy_chk:
                                blank_col = legacy_chk.get("blank_column")
                                date_col_chk = legacy_chk.get("date_column")
                                check_plan = legacy_chk.get("check_planning_months", True)
                                legacy_demands = []
                                for r_idx in rows:
                                    if r_idx in metadata_rows:
                                        continue
                                    if blank_col and _identifier(ws.cell(r_idx, blank_col).value):
                                        continue
                                    has_plan = (
                                        any(ws.cell(r_idx, col).value is not None for col in contract.col_to_month)
                                        if check_plan
                                        else False
                                    )
                                    has_date = (
                                        ws.cell(r_idx, date_col_chk).value is not None
                                        if date_col_chk and original_columns >= date_col_chk
                                        else False
                                    )
                                    if has_plan or has_date:
                                        legacy_demands.append(r_idx)
                                if legacy_demands:
                                    event(
                                        s,
                                        "AMBIGUOUS_LEGACY_REQUIREMENT",
                                        r,
                                        confidence="REVIEW",
                                        reason="Anonymous planner demand exists; source identity cannot be established",
                                        candidate_original_rows=legacy_demands,
                                    )
                                    continue

                        ins_status = matched_rule.get(
                            "insert_status",
                            "INSERTED_RECORD_ROW",
                        )
                        schedule_insert(
                            s, anchor, row_values, ins_status, r, matched_rule.get("date_number_format", "dd.mm.yyyy")
                        )
                        insertions[anchor][-1]["event"]["chosen_quantity"] = qty

            def moved(row_idx: int) -> int:
                return row_idx + sum(len(items) for anch, items in insertions.items() if anch < row_idx)

            for anchor, items in sorted(insertions.items(), reverse=True):
                style_row = contract.data_start_row if populate_template else anchor
                styles = [copy.copy(ws.cell(style_row, col)._style) for col in range(1, ws.max_column + 1)]
                ws.insert_rows(anchor + 1, len(items))
                for offset, item in enumerate(items, 1):
                    new_r = anchor + offset
                    for col_idx, st in enumerate(styles, 1):
                        ws.cell(new_r, col_idx)._style = copy.copy(st)
                    for col_idx, val in item["values"].items():
                        ws.cell(new_r, col_idx).value = val
                        if isinstance(val, str):
                            ws.cell(new_r, col_idx).data_type = "s"
                        if isinstance(val, (datetime.date, datetime.datetime)):
                            ws.cell(new_r, col_idx).number_format = item["date_format"]
                    final_r = moved(anchor) + offset
                    item["event"]["row"] = final_r
                    item["row"] = final_r
                    if "record_id" in item["event"]:
                        metadata[item["event"]["record_id"]] = {
                            "row": final_r,
                            "final": True,
                            "values": {str(c): _plain(v) for c, v in item["values"].items()},
                        }

            ws.row_dimensions.clear()
            for r_idx, dim in original_dimensions.items():
                dim.index = moved(r_idx)
                ws.row_dimensions[moved(r_idx)] = dim
            for anchor, items in insertions.items():
                for item in items:
                    style_row = contract.data_start_row if populate_template else anchor
                    if style_row in original_dimensions:
                        dim = copy.copy(original_dimensions[style_row])
                        dim.index = item["row"]
                        ws.row_dimensions[item["row"]] = dim

            if freeze:
                freeze_col, freeze_row = coordinate_from_string(freeze)
                ws.freeze_panes = f"{freeze_col}{moved(freeze_row)}"

            for item in events:
                if "original_row" in item:
                    item["row"] = moved(item["original_row"])
            for entry in metadata.values():
                if not entry.pop("final", False):
                    entry["row"] = moved(entry["row"])

            if metadata_sheet_name in book:
                del book[metadata_sheet_name]
            meta = book.create_sheet(metadata_sheet_name)
            meta.append(DEFAULT_METADATA_HEADER)
            current_hash = self.profile_hash
            for ident, entry in sorted(metadata.items()):
                meta.append(
                    [1, ws.title, entry["row"], ident, json.dumps(entry["values"], ensure_ascii=False), current_hash]
                )
            meta.sheet_state = "hidden"

            summary = {
                "mode": "populate_template" if populate_template else "update_existing",
                "sheet_name": ws.title,
                "baseline_rows": original_rows,
                "row_mapping": {str(r): moved(r) for r in range(1, original_rows + 1)},
                "allowed_changes": changes,
                "allowed_style_changes": [
                    {"row": r, "column": c, **info} for (r, c), info in allowed_style_changes.items()
                ],
                "events": events,
                "inserted_rows": [
                    {"row": it["row"], "values": {str(c): _plain(v) for c, v in it["values"].items()}}
                    for items in insertions.values()
                    for it in items
                ],
                "relevant_sections_count": len(relevant),
                "excluded_sections_count": len(sections) - len(relevant),
                "relevant_materials_count": len(relevant),
                "excluded_materials_count": len(sections) - len(relevant),
                "total_rows_inserted": sum(map(len, insertions.values())),
                "out_of_horizon_records_count": sum(e["status"] == "OUTSIDE_PLANNING_HORIZON" for e in events),
                "status_counts": dict(Counter(e["status"] for e in events)),
                "record_disposition_counts": dict(Counter(e["status"] for e in events if "record_id" in e)),
                "normalized_record_count": sum(
                    len(getattr(s, "records", getattr(s, "all_records", []))) for s in sections
                ),
                "accounted_record_count": sum("record_id" in e for e in events),
                "horizon_months": list(contract.month_to_col),
                "metadata_sheet_name": metadata_sheet_name,
            }
            if summary["normalized_record_count"] != summary["accounted_record_count"]:
                raise ValueError("Each normalized source record must have exactly one final disposition.")

            output_path.parent.mkdir(parents=True, exist_ok=True)
            descriptor, name = tempfile.mkstemp(suffix=".xlsx", dir=output_path.parent)
            os.close(descriptor)
            temporary = Path(name)
            book.save(temporary)

            comparator = WorkbookPreservationComparator()
            comparison = comparator.compare_workbooks(self.template_path, temporary, summary)
            if comparison["status"] != "PASS":
                raise ValueError(f"Workbook preservation verification failed: {comparison['issues'][:5]}")

            os.replace(temporary, output_path)
            return summary
        finally:
            book.close()
            if temporary is not None and temporary.exists():
                temporary.unlink()

    def _eval_spec(self, spec: Any, section: Any, record: Any) -> Any:
        raw = record if isinstance(record, dict) else getattr(record, "raw", None)
        return _resolve_value_spec(spec, section, raw)

    def _template_section_sort_key(self, section: Any) -> tuple:
        values = []
        for spec in self.policy["template_section_sort"]:
            value = self._eval_spec(spec, section, None)
            if value in (None, ""):
                values.append((2, ""))
            elif isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
                values.append((0, value))
            elif isinstance(value, str):
                values.append((1, value.casefold()))
            else:
                raise ValueError("Template section sort values must be finite numbers or text.")
        return tuple(values)

    def _read_metadata(
        self, book: Any, ws: Any, metadata_sheet_name: str, contract: WorkbookLayoutContract
    ) -> dict[str, Any]:
        if metadata_sheet_name not in book:
            return {}
        sheet = book[metadata_sheet_name]
        header_vals = tuple(cell.value for cell in sheet[1])
        if header_vals != DEFAULT_METADATA_HEADER:
            raise ValueError(f"Unrecognized provenance sheet '{metadata_sheet_name}'; refusing to replace it.")
        res = {}
        seen_rows = set()
        for row_data in sheet.iter_rows(min_row=2, values_only=True):
            if not row_data or all(value is None for value in row_data):
                continue
            version = row_data[0]
            title = row_data[1] if len(row_data) > 1 else None
            row = row_data[2] if len(row_data) > 2 else None
            identity = row_data[3] if len(row_data) > 3 else None
            values_str = row_data[4] if len(row_data) > 4 else None

            if (
                type(version) is not int
                or version != 1
                or len(row_data) != len(DEFAULT_METADATA_HEADER)
                or row_data[5] != self.profile_hash
                or title != ws.title
                or not isinstance(identity, str)
                or re.fullmatch(r"[a-f0-9]{64}", identity) is None
                or identity in res
                or isinstance(row, bool)
                or not isinstance(row, int)
                or not contract.data_start_row <= row <= ws.max_row
                or row in seen_rows
            ):
                raise ValueError("Invalid or conflicting provenance metadata.")
            seen_rows.add(row)

            try:
                values_dict = json.loads(values_str)
                if not isinstance(values_dict, dict) or not values_dict:
                    raise ValueError("Provenance written_values must be a JSON object.")
                for col_k, col_v in values_dict.items():
                    col_int = int(col_k)
                    if str(col_int) != col_k or not (1 <= col_int <= ws.max_column):
                        raise ValueError(f"Column {col_k} out of bounds in provenance.")
                    if isinstance(col_v, float) and not math.isfinite(col_v):
                        raise ValueError("Non-finite value in provenance.")
                    if col_v is not None and not isinstance(col_v, (str, int, float)):
                        raise ValueError(f"Unsupported value type for column {col_k} in provenance.")
            except (ValueError, TypeError) as exc:
                raise ValueError("Provenance written_values JSON is malformed.") from exc

            res[identity] = {"row": row, "values": values_dict}
        return res


class WorkbookPreservationComparator:
    """Compare baseline cells against generated output using original-to-generated row mapping."""

    def compare_workbooks(
        self, baseline_path: Path, generated_path: Path, validation_summary: dict[str, Any]
    ) -> dict[str, Any]:
        base = openpyxl.load_workbook(baseline_path, data_only=False)
        generated = openpyxl.load_workbook(generated_path, data_only=False)
        issues = []
        checked = 0
        try:
            title = validation_summary["sheet_name"]
            mapping = validation_summary["row_mapping"]
            changes = {(item["original_row"], item["column"]): item for item in validation_summary["allowed_changes"]}
            allowed_styles = {
                (item["original_row"], item["column"]): item
                for item in validation_summary.get("allowed_style_changes", [])
            }
            meta_sheet = validation_summary.get("metadata_sheet_name", DEFAULT_METADATA_SHEET)

            if base.defined_names != generated.defined_names:
                issues.append({"type": "defined_names"})
            if [n for n in generated.sheetnames if n != meta_sheet] != [n for n in base.sheetnames if n != meta_sheet]:
                issues.append({"type": "sheet_names"})

            for original in base:
                if original.title == meta_sheet:
                    continue
                final = generated[original.title]
                for row in original:
                    for cell in row:
                        target_row = mapping[str(cell.row)] if original.title == title else cell.row
                        actual = final.cell(target_row, cell.column)
                        change = changes.get((cell.row, cell.column)) if original.title == title else None
                        expected = change["new_value"] if change else _plain(cell.value)
                        value_matches = (_plain(actual.value) == expected) if change else (actual.value == cell.value)
                        if not value_matches:
                            issues.append(
                                {
                                    "type": "cell_value",
                                    "sheet": original.title,
                                    "original": cell.coordinate,
                                    "generated": actual.coordinate,
                                }
                            )
                        if (cell.has_style or actual.has_style) and cell._style != actual._style:
                            style_change = (
                                allowed_styles.get((cell.row, cell.column)) if original.title == title else None
                            )
                            if not _same_style(cell, actual, style_change):
                                issues.append({"type": "cell_style", "original": cell.coordinate})
                        if cell.comment != actual.comment or cell.hyperlink != actual.hyperlink:
                            issues.append({"type": "comment_or_link", "original": cell.coordinate})
                        checked += 1

                for key, dimension in original.column_dimensions.items():
                    if dict(dimension) != dict(final.column_dimensions[key]):
                        issues.append({"type": "column_dimension", "column": key})
                for key, dimension in original.row_dimensions.items():
                    target = mapping[str(key)] if original.title == title and str(key) in mapping else key
                    expected = dict(dimension)
                    expected.pop("r", None)
                    expected.pop("s", None)
                    actual = dict(final.row_dimensions[target])
                    actual.pop("r", None)
                    actual.pop("s", None)
                    if expected != actual or not _same_style(dimension, final.row_dimensions[target]):
                        issues.append({"type": "row_dimension", "row": key, "expected": expected, "actual": actual})

                expected_freeze = original.freeze_panes
                if expected_freeze and original.title == title:
                    freeze_col, freeze_row = coordinate_from_string(expected_freeze)
                    expected_freeze = f"{freeze_col}{mapping.get(str(freeze_row), freeze_row)}"
                if final.freeze_panes != expected_freeze:
                    issues.append({"type": "freeze_panes"})

                for attr in (
                    "sheet_format",
                    "sheet_properties",
                    "sheet_state",
                    "page_setup",
                    "page_margins",
                    "print_options",
                ):
                    if getattr(original, attr) != getattr(final, attr):
                        issues.append({"type": attr, "sheet": original.title})

            ws = generated[title]
            for item in validation_summary["inserted_rows"]:
                for col, expected in item["values"].items():
                    if _plain(ws.cell(item["row"], int(col)).value) != expected:
                        issues.append({"type": "inserted_record_value", "row": item["row"], "column": col})

            # A header-only template may end before data_start_row. Its first
            # generated rows also establish the previously unused data area.
            expected_rows = max(
                base[title].max_row + validation_summary["total_rows_inserted"],
                max((row["row"] for row in validation_summary["inserted_rows"]), default=0),
            )
            if ws.max_row != expected_rows:
                issues.append({"type": "row_count", "expected": expected_rows, "actual": ws.max_row})

            return {
                "status": "PASS" if not issues else "FAIL",
                "issues": issues,
                "baseline_cells_checked": checked,
                "baseline_rows": base[title].max_row,
                "generated_rows": ws.max_row,
                "inserted_rows_checked": len(validation_summary["inserted_rows"]),
                "reference_workbook_status": "NOT_PROVIDED",
                "reference_row_match_metrics": None,
                "review_required": any(e.get("confidence") == "REVIEW" for e in validation_summary["events"]),
            }
        finally:
            base.close()
            generated.close()


def compare_reference_workbook(
    reference_path: Path,
    generated_path: Path,
    sheet_name: str,
    *,
    contract: WorkbookLayoutContract | None = None,
    expected_headers: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Measure a supplied row fixture, or explicitly identify an unsupported/prose reference."""
    reference = openpyxl.load_workbook(reference_path, read_only=True, data_only=False)
    generated = openpyxl.load_workbook(generated_path, read_only=True, data_only=False)
    try:
        hdr_row = contract.header_row if contract else 1
        sec_col = contract.section_key_column if contract else 1
        exp_hdrs = expected_headers or (getattr(contract, "expected_headers", None)) or {}

        candidates = []
        for sheet in reference:
            if exp_hdrs:
                matches_all = True
                for col_str, accepted in exp_hdrs.items():
                    col_int = int(col_str)
                    actual_val = str(sheet.cell(hdr_row, col_int).value or "").strip()
                    aliases = accepted if isinstance(accepted, (list, tuple)) else [accepted]
                    if not any(_canonical_header(a) in _canonical_header(actual_val) for a in aliases):
                        matches_all = False
                        break
                if matches_all:
                    candidates.append(sheet)
            else:
                cell_val = str(sheet.cell(hdr_row, sec_col).value or "").strip()
                if cell_val:
                    candidates.append(sheet)

        if len(candidates) != 1:
            return {
                "reference_workbook_status": "PROSE_OR_UNSUPPORTED_LAYOUT; corrected-row fixture unavailable",
                "reference_row_match_metrics": None,
            }
        expected_sheet = candidates[0]
        if sheet_name not in generated.sheetnames:
            return {
                "reference_workbook_status": f"UNSUPPORTED_LAYOUT; target sheet '{sheet_name}' missing",
                "reference_row_match_metrics": None,
            }
        actual_sheet = generated[sheet_name]
        if expected_sheet.max_row > 100_000:
            raise ValueError("Reference workbook exceeds the supported comparison row limit.")

        min_r = contract.data_start_row if contract else 2
        sec_col_0based = sec_col - 1

        def extract_records(sheet: Any) -> list[tuple[Any, ...]]:
            result = []
            for row in sheet.iter_rows(min_row=min_r, values_only=True):
                if len(row) > sec_col_0based and row[sec_col_0based] is not None:
                    result.append(tuple(_plain(v) for v in row))
            return result

        expected_records = extract_records(expected_sheet)
        actual_records = extract_records(actual_sheet)
        exp_counts, act_counts = Counter(expected_records), Counter(actual_records)
        matching = sum((exp_counts & act_counts).values())

        missing_keys = sorted(
            {r[sec_col_0based] for r in expected_records} - {r[sec_col_0based] for r in actual_records}
        )
        unexpected_keys = sorted(
            {r[sec_col_0based] for r in actual_records} - {r[sec_col_0based] for r in expected_records}
        )

        return {
            "reference_workbook_status": "PLANNING_ROW_FIXTURE",
            "reference_row_match_metrics": {
                "matching_rows": matching,
                "missing_rows": len(expected_records) - matching,
                "unexpected_rows": len(actual_records) - matching,
                "missing_section_keys": missing_keys,
                "unexpected_section_keys": unexpected_keys,
            },
        }
    finally:
        reference.close()
        generated.close()
