"""Selected PDFs replace source data in either time direction; planner values survive separately."""

import copy
import datetime

import openpyxl
import pytest
from openpyxl.comments import Comment

from shogun.services.sectioned_workbook_pipeline import SectionedLayoutParser
from shogun.services.sectioned_workbook_updater import SafeWorkbookUpdater, WorkbookPreservationComparator
from shogun.services.workbook_master_refresh import PLANNER_SHEET
from tests.test_workbook_master_maintenance import master_setup as master_fixture
from tests.test_workbook_master_maintenance import section

master_setup = master_fixture


def generate(seed, profile, text, output):
    sections = SectionedLayoutParser(profile).parse_pages(list(enumerate(text.split("\f"), 1)), "selected.pdf")
    return SafeWorkbookUpdater(seed, profile).execute(
        sections, output, options={"master_refresh_mode": "selected_pdfs"},
    )


def values(path):
    book = openpyxl.load_workbook(path)
    result = list(book["Plan"].values)
    book.close()
    return result


@pytest.mark.parametrize("first_month,second_month", [("2026/08", "2026/10"), ("2026/10", "2026/08")])
def test_selected_sources_replace_every_pdf_cell_both_directions(master_setup, tmp_path, first_month, second_month):
    seed, profile = master_setup
    old, current, expected = (tmp_path / name for name in ("old.xlsx", "current.xlsx", "expected.xlsx"))
    generate(seed, profile, section(report=first_month) + "\f" + section("OLD", report=first_month, records=""), old)
    selected = section("NEW", report=second_month, stock=7, records=(
        f"JOB Q1 18 {second_month} {second_month.replace('/', '-')}-12\n"
        f"DEM Q2 9 {second_month} {second_month.replace('/', '-')}-19\n"
    ))
    summary = generate(old, profile, selected, current)
    generate(seed, profile, selected, expected)
    assert values(current) == values(expected)
    assert not any("OLD" in row or "P1" in row or "D1" in row for row in values(current))
    assert summary["master_refresh_mode"] == "selected_pdfs"
    assert WorkbookPreservationComparator().compare_workbooks(old, current, summary)["status"] == "PASS"


def test_notes_and_absent_rows_survive_round_trip_without_pdf_overrides(master_setup, tmp_path):
    seed, profile = master_setup
    august, october, returned = (tmp_path / name for name in ("august.xlsx", "october.xlsx", "returned.xlsx"))
    generate(seed, profile, section(), august)
    book = openpyxl.load_workbook(august)
    sheet = book["Plan"]
    row = next(row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 5).value == "P1")
    sheet.cell(row, 7, "Planner customer")
    sheet.cell(row, 8, "Keep my note")
    sheet.cell(row, 9, "Avo 17")
    sheet.cell(row, 8).comment = Comment("Also keep this annotation", "Planner")
    sheet.cell(row, 8).font = openpyxl.styles.Font(color="FF0000", bold=True)
    sheet.cell(row, 6, 999)  # This may be retained, but must not override the selected PDF quantity.
    sheet.cell(row, 11, 77)  # A manual month quantity must not contaminate a different period.
    extra = book.create_sheet("Planner notebook")
    extra["A1"] = "Keep unrelated worksheet"
    book.save(august)
    book.close()
    generate(august, profile, section("B", report="2026/10", records="JOB Q1 8 2026/10 2026-10-12\n"), october)
    book = openpyxl.load_workbook(october)
    assert all(row[1] != "A" for row in list(book["Plan"].values)[2:])
    archived = [row[3] for row in book[PLANNER_SHEET].iter_rows(min_row=2, values_only=True)]
    assert all(value in archived for value in ("Planner customer", "Keep my note", "Avo 17", 999, 77))
    assert book["Planner notebook"]["A1"].value == "Keep unrelated worksheet"
    book.close()
    summary = generate(october, profile, section(), returned)
    book = openpyxl.load_workbook(returned)
    sheet = book["Plan"]
    row = next(row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 5).value == "P1")
    assert [sheet.cell(row, col).value for col in (6, 7, 8, 9)] == [30, "Planner customer", "Keep my note", "Avo 17"]
    assert sheet.cell(row, 8).comment.text == "Also keep this annotation"
    assert sheet.cell(row, 8).font.color.rgb == "00FF0000"
    assert sheet.cell(row, 11).value is None
    assert summary["status_counts"]["PLANNER_VALUE_ARCHIVED"] == 2
    book.close()


def test_cleared_active_planner_value_is_not_resurrected(master_setup, tmp_path):
    seed, profile = master_setup
    first, second, third = (tmp_path / name for name in ("one.xlsx", "two.xlsx", "three.xlsx"))
    generate(seed, profile, section(), first)
    book = openpyxl.load_workbook(first)
    book["Plan"]["G3"] = "Delete later"
    book.save(first)
    book.close()
    generate(first, profile, section(), second)
    book = openpyxl.load_workbook(second)
    book["Plan"]["G3"] = None
    book.save(second)
    book.close()
    generate(second, profile, section(), third)
    book = openpyxl.load_workbook(third)
    assert book["Plan"]["G3"].value is None
    assert book[PLANNER_SHEET].max_row == 1
    book.close()


def test_changed_pdf_wins_over_old_manual_quantity_and_notes_stay_with_identity(master_setup, tmp_path):
    seed, profile = master_setup
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    generate(seed, profile, section(), first)
    book = openpyxl.load_workbook(first)
    row = next(row for row in range(3, book["Plan"].max_row + 1) if book["Plan"].cell(row, 5).value == "P1")
    book["Plan"].cell(row, 7, "Correct customer")
    book["Plan"].cell(row, 6, 999)
    book.save(first)
    book.close()
    generate(first, profile, section(report="2026/10", stock=0, records="JOB P1 42 2026/10 2026-10-19\n"), second)
    book = openpyxl.load_workbook(second)
    sheet = book["Plan"]
    assert sheet.max_row == 4  # Only current source row and preserved blank template row.
    assert sheet["F3"].value == 42
    assert sheet["G3"].value == "Correct customer"
    assert sheet["J3"].value == datetime.datetime(2026, 10, 19)
    assert 999 in [row[3] for row in book[PLANNER_SHEET].iter_rows(min_row=2, values_only=True)]
    book.close()


def test_mutated_output_is_rejected_by_comparator(master_setup, tmp_path):
    seed, profile = master_setup
    output = tmp_path / "output.xlsx"
    summary = generate(seed, profile, section(), output)
    book = openpyxl.load_workbook(output)
    book["Plan"]["G3"] = "unexpected change"
    book.save(output)
    book.close()
    assert WorkbookPreservationComparator().compare_workbooks(seed, output, summary)["status"] == "FAIL"


def test_legacy_master_migrates_without_changing_profile(master_setup, tmp_path):
    seed, profile = master_setup
    from tests.test_workbook_master_maintenance import run_master

    legacy, output = tmp_path / "legacy.xlsx", tmp_path / "selected.xlsx"
    original_profile = copy.deepcopy(profile)
    run_master(seed, profile, section(report="2026/10"), legacy)
    generate(legacy, profile, section(report="2026/08"), output)
    assert profile == original_profile
    assert values(output)[0][10] == datetime.datetime(2026, 8, 1)


def test_unassigned_note_is_archived_and_archive_survives_identical_reruns(master_setup, tmp_path):
    seed, profile = master_setup
    first, second, third = (tmp_path / name for name in ("one.xlsx", "two.xlsx", "three.xlsx"))
    generate(seed, profile, section(), first)
    book = openpyxl.load_workbook(first)
    book["Plan"].cell(book["Plan"].max_row, 8, "Unassigned planner note")
    book.save(first)
    book.close()
    generate(first, profile, section(report="2026/10"), second)
    generate(second, profile, section(report="2026/10"), third)
    assert values(second) == values(third)
    book = openpyxl.load_workbook(third)
    assert "Unassigned planner note" in [row[3] for row in book[PLANNER_SHEET].iter_rows(min_row=2, values_only=True)]
    book.close()


def test_archive_collision_fails_without_publication(master_setup, tmp_path):
    seed, profile = master_setup
    output = tmp_path / "output.xlsx"
    book = openpyxl.load_workbook(seed)
    book.create_sheet(PLANNER_SHEET)["A1"] = "Not a Shogun archive"
    book.save(seed)
    book.close()
    with pytest.raises(ValueError, match="not a recognized"):
        generate(seed, profile, section(), output)
    assert not output.exists()


def test_changed_row_structure_fails_without_publication(master_setup, tmp_path):
    seed, profile = master_setup
    first, output = tmp_path / "first.xlsx", tmp_path / "output.xlsx"
    generate(seed, profile, section(), first)
    book = openpyxl.load_workbook(first)
    book["Plan"].insert_rows(3)
    book.save(first)
    book.close()
    before = first.read_bytes()
    with pytest.raises(ValueError, match="metadata|ownership|provenance"):
        generate(first, profile, section(report="2026/10"), output)
    assert first.read_bytes() == before
    assert not output.exists()
