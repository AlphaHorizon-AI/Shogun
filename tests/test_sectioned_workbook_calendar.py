"""A reusable empty planning template follows the current source report month."""

import copy
import datetime
import hashlib

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from shogun.services.sectioned_workbook_pipeline import SectionedLayoutParser
from shogun.services.sectioned_workbook_updater import (
    SafeWorkbookUpdater,
    WorkbookPreservationComparator,
    validate_workbook_update_profile,
)


@pytest.fixture
def calendar_population(tmp_path):
    template = tmp_path / "empty.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Schedule"
    sheet.append(["Description", "Equipment", "Job", "Hours", "Due date", "Note"])
    for offset in range(12):
        cell = sheet.cell(1, 7 + offset, datetime.datetime(2020, 1 + offset, 1))
        cell.number_format = "mmm yyyy"
        cell.fill = PatternFill("solid", fgColor="AADDCC")
    sheet.cell(3, 2).fill = PatternFill("solid", fgColor="AADDCC")
    sheet.freeze_panes = "G3"
    book.save(template)
    book.close()
    profile = {
        "id": "equipment_calendar",
        "adapter": "sectioned_record_matrix_v1",
        "parameters": {
            "section_pattern": r"(?m)^Equipment: (?P<id>\S+)$",
            "section_key_group": "id",
            "section_fields": [
                {"target": "report_year", "pattern": r"Report: (?P<value>\d{4})/", "group": "value"},
                {"target": "report_month", "pattern": r"Report: \d{4}/(?P<value>\d{2})", "group": "value"},
            ],
            "record_pattern": r"(?m)^(?P<job>JOB-\d+) (?P<hours>\d+) (?P<date>\d{4}-\d{2}-\d{2})$",
            "extraction": {"candidate_line_pattern": r"^JOB-"},
            "row_rules": [{"kind": "record", "columns": {"0": {"group": "job"}}}],
            "workbook_update": {
                "mode": "populate_template",
                "sheet_name": "Schedule",
                "section_key_column": 2,
                "data_start_row": 3,
                "expected_headers": {"2": "Equipment"},
                "planning_start_column": 7,
                "require_planning_months": True,
                "template_planning_months": {
                    "start_month_spec": {
                        "join": {
                            "values": [{"field": "report_year"}, {"field": "report_month"}],
                            "separator": "/",
                            "require_all": True,
                        },
                    },
                    "count": 12,
                },
                "source_identity_fields": [{"section_key": True}, {"group": "job"}],
                "record_rules": [{
                    "id": "job",
                    "match": {},
                    "quantity_spec": {"group": "hours", "value_type": "number"},
                    "date_spec": {"group": "date", "value_type": "iso_date"},
                    "planning_month_quantity": True,
                    "require_date_in_planning_horizon": True,
                    "insert_row_values": {"2": {"section_key": True}, "3": {"group": "job"}},
                }],
            },
        },
    }
    return template, profile


def parse_sections(profile, text):
    return SectionedLayoutParser(profile).parse_pages(
        list(enumerate(text.split("\f"), 1)), source_file="current.pdf",
    )


def source_text(month):
    return f"Equipment: UNIT-1\nReport: {month}\nJOB-1 7 {month.replace('/', '-')}-06\n"


@pytest.mark.parametrize("reference,last_within,first_after", [
    ("2026-10-19", "2027-01-19", "2027-01-20"),
    ("2024-01-31", "2024-04-30", "2024-05-01"),
    ("2024-11-30", "2025-02-28", "2025-03-01"),
])
def test_inserted_rows_date_band_uses_inclusive_calendar_month_cutoff(
    calendar_population, tmp_path, reference, last_within, first_after,
):
    template, profile = calendar_population
    rule = profile["parameters"]["workbook_update"]["record_rules"][0]
    rule["inserted_row_date_band_fill"] = {
        "reference_date_spec": {"literal": reference, "value_type": "iso_date"},
        "date_spec": {"group": "date", "value_type": "iso_date"},
        "months": 3,
        "columns": [1, 5],
        "within_color": "FFDDF3DE",
        "after_color": "FFDCEBFA",
    }
    year, month = reference[:7].split("-")
    sections = parse_sections(profile, f"Equipment: UNIT-1\nReport: {year}/{month}\n"
                              f"JOB-1 7 {last_within}\nJOB-2 8 {first_after}\n")
    output = tmp_path / "bands.xlsx"
    summary = SafeWorkbookUpdater(template, profile).execute(sections, output)
    assert WorkbookPreservationComparator().compare_workbooks(template, output, summary)["status"] == "PASS"
    book = openpyxl.load_workbook(output)
    try:
        sheet = book["Schedule"]
        jobs = {sheet.cell(row, 3).value: row for row in range(3, sheet.max_row + 1)}
        for job, color in [("JOB-1", "FFDDF3DE"), ("JOB-2", "FFDCEBFA")]:
            row = jobs[job]
            assert sheet.cell(row, 1).fill.fgColor.rgb == color
            assert sheet.cell(row, 5).fill.fgColor.rgb == color
            assert sheet.cell(row, 4).fill.fgColor.rgb != color
    finally:
        book.close()


def test_inserted_row_date_band_rejects_missing_dates(calendar_population, tmp_path):
    template, profile = calendar_population
    rule = profile["parameters"]["workbook_update"]["record_rules"][0]
    rule["inserted_row_date_band_fill"] = {
        "reference_date_spec": {"field": "missing"}, "date_spec": {"group": "date", "value_type": "iso_date"},
        "months": 3, "columns": [1], "within_color": "FFDDF3DE", "after_color": "FFDCEBFA",
    }
    output = tmp_path / "invalid-date-band.xlsx"
    with pytest.raises(ValueError, match="requires valid source dates"):
        SafeWorkbookUpdater(template, profile).execute(parse_sections(profile, source_text("2026/10")), output)
    assert not output.exists()


@pytest.mark.parametrize("changes", [
    {"months": True}, {"months": -1}, {"months": 121}, {"columns": []}, {"columns": [1, 1]},
    {"within_color": "red"}, {"unexpected": "value"},
])
def test_inserted_row_date_band_schema_rejects_invalid_settings(calendar_population, changes):
    _, profile = calendar_population
    band = {
        "reference_date_spec": {"literal": "2026-10-19", "value_type": "iso_date"},
        "date_spec": {"group": "date", "value_type": "iso_date"}, "months": 3,
        "columns": [1], "within_color": "FFDDF3DE", "after_color": "FFDCEBFA",
    }
    band.update(changes)
    profile["parameters"]["workbook_update"]["record_rules"][0]["inserted_row_date_band_fill"] = band
    with pytest.raises(ValueError, match="inserted_row_date_band_fill"):
        validate_workbook_update_profile(profile)


def test_inserted_row_date_band_does_not_restyle_existing_workbooks(calendar_population):
    _, profile = calendar_population
    policy = profile["parameters"]["workbook_update"]
    del policy["template_planning_months"]
    policy["mode"] = "update_existing"
    policy["record_rules"][0]["inserted_row_date_band_fill"] = {
        "reference_date_spec": {"literal": "2026-10-19", "value_type": "iso_date"},
        "date_spec": {"group": "date", "value_type": "iso_date"}, "months": 3,
        "columns": [1], "within_color": "FFDDF3DE", "after_color": "FFDCEBFA",
    }
    with pytest.raises(ValueError, match="requires populate_template"):
        validate_workbook_update_profile(profile)


def test_same_template_rolls_to_two_source_report_months(calendar_population, tmp_path):
    template, profile = calendar_population
    original_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    for month in ("2026/07", "2026/10"):
        output = tmp_path / f"result_{month.replace('/', '_')}.xlsx"
        sections = parse_sections(profile, source_text(month))
        summary = SafeWorkbookUpdater(template, profile).execute(sections, output)
        comparison = WorkbookPreservationComparator().compare_workbooks(template, output, summary)
        assert comparison["status"] == "PASS"
        assert summary["horizon_months"][0] == month
        assert summary["record_disposition_counts"] == {"INSERTED_RECORD_ROW": 1}
        assert len(summary["allowed_changes"]) == 12
        book = openpyxl.load_workbook(output)
        try:
            sheet = book["Schedule"]
            assert sheet["G1"].value.strftime("%Y/%m") == month
            assert sheet["G1"].number_format == "mmm yyyy"
            assert sheet["G1"].fill.fgColor.rgb == "00AADDCC"
            assert sheet.freeze_panes == "G3"
            assert sheet["G4"].value == 7
            assert all(sheet.cell(4, column).value is None for column in range(8, 19))
        finally:
            book.close()
    assert hashlib.sha256(template.read_bytes()).hexdigest() == original_hash


def test_rolling_calendar_crosses_year_boundary(calendar_population, tmp_path):
    template, profile = calendar_population
    output = tmp_path / "rollover.xlsx"
    sections = parse_sections(profile, source_text("2026/11"))
    summary = SafeWorkbookUpdater(template, profile).execute(sections, output)
    assert summary["horizon_months"] == [
        "2026/11", "2026/12", "2027/01", "2027/02", "2027/03", "2027/04",
        "2027/05", "2027/06", "2027/07", "2027/08", "2027/09", "2027/10",
    ]


@pytest.mark.parametrize("report,error", [
    ("Report: 2026/10\n", "conflicting source months"),
    ("", "valid source month from every section"),
    ("Report: 2026/13\n", "valid source month from every section"),
])
def test_every_section_must_agree_even_without_records_or_when_excluded(
    calendar_population, tmp_path, report, error,
):
    template, profile = calendar_population
    profile["parameters"]["section_selection"] = {"field": "id", "operator": "equals", "value": "UNIT-1"}
    sections = parse_sections(profile, source_text("2026/07") + f"\fEquipment: EXCLUDED\n{report}")
    before = template.read_bytes()
    output = tmp_path / "invalid.xlsx"
    with pytest.raises(ValueError, match=error):
        SafeWorkbookUpdater(template, profile).execute(sections, output)
    assert not output.exists()
    assert template.read_bytes() == before


@pytest.mark.parametrize("change,error", [
    ({"mode": "update_existing"}, "requires populate_template mode"),
    ({"planning_start_column": None}, "requires planning_start_column"),
    ({"planning_start_column": 16380}, "Excel column limit"),
    ({"template_planning_months": {"count": 12}}, "start_month_spec and count"),
    ({"template_planning_months": {"start_month_spec": "month", "count": 12}}, "must be a dictionary"),
    *[
        ({"template_planning_months": {"start_month_spec": {"field": "month"}, "count": count}},
         "integer from one to 120")
        for count in (True, 0, 121, 12.0)
    ],
])
def test_rolling_calendar_profile_validation(calendar_population, change, error):
    _, original = calendar_population
    profile = copy.deepcopy(original)
    profile["parameters"]["workbook_update"].update(change)
    with pytest.raises(ValueError, match=error):
        validate_workbook_update_profile(profile)


@pytest.mark.parametrize("change", ["non_date", "missing", "extra", "short_count"])
def test_calendar_refuses_to_replace_or_shrink_other_header_layouts(calendar_population, tmp_path, change):
    template, profile = calendar_population
    book = openpyxl.load_workbook(template)
    sheet = book["Schedule"]
    if change == "non_date":
        sheet["H1"] = "Overdue"
        profile["parameters"]["workbook_update"]["backlog_headers"] = ["Overdue"]
    elif change == "missing":
        sheet["H1"] = None
    elif change == "extra":
        sheet["S1"] = datetime.datetime(2021, 1, 1)
    else:
        profile["parameters"]["workbook_update"]["template_planning_months"]["count"] = 11
    book.save(template)
    book.close()
    before = template.read_bytes()
    output = tmp_path / "invalid.xlsx"
    with pytest.raises(ValueError, match="exactly count contiguous dated header cells"):
        SafeWorkbookUpdater(template, profile).execute(parse_sections(profile, source_text("2026/07")), output)
    assert not output.exists()
    assert template.read_bytes() == before


def test_without_rolling_configuration_preserves_static_months(calendar_population, tmp_path):
    template, profile = calendar_population
    del profile["parameters"]["workbook_update"]["template_planning_months"]
    output = tmp_path / "static.xlsx"
    summary = SafeWorkbookUpdater(template, profile).execute(parse_sections(profile, source_text("2026/07")), output)
    assert summary["horizon_months"] == [f"2020/{month:02}" for month in range(1, 13)]
    assert summary["allowed_changes"] == []
    assert summary["record_disposition_counts"] == {"OUTSIDE_PLANNING_HORIZON": 1}


def test_calendar_requires_source_sections(calendar_population, tmp_path):
    template, profile = calendar_population
    output = tmp_path / "no_sources.xlsx"
    with pytest.raises(ValueError, match="source month from every section"):
        SafeWorkbookUpdater(template, profile).execute([], output)
    assert not output.exists()


def test_rolling_calendar_rejects_date_overflow_before_publication(calendar_population, tmp_path):
    template, profile = calendar_population
    sections = parse_sections(profile, "Equipment: UNIT-1\nReport: 9999/12\n")
    output = tmp_path / "overflow.xlsx"
    with pytest.raises(ValueError, match="supported calendar year range"):
        SafeWorkbookUpdater(template, profile).execute(sections, output)
    assert not output.exists()


def test_update_existing_still_moves_freeze_with_preserved_rows(calendar_population, tmp_path):
    template, profile = calendar_population
    policy = profile["parameters"]["workbook_update"]
    policy["mode"] = "update_existing"
    del policy["template_planning_months"]
    book = openpyxl.load_workbook(template)
    sheet = book["Schedule"]
    sheet["B3"] = "UNIT-1"
    sheet["B4"] = "UNIT-2"
    sheet.freeze_panes = "G4"
    book.save(template)
    book.close()
    output = tmp_path / "updated.xlsx"
    sections = parse_sections(profile, source_text("2020/01"))
    summary = SafeWorkbookUpdater(template, profile).execute(sections, output)
    assert summary["total_rows_inserted"] == 1
    comparison = WorkbookPreservationComparator().compare_workbooks(template, output, summary)
    assert comparison["status"] == "PASS"
    book = openpyxl.load_workbook(output)
    try:
        assert book["Schedule"].freeze_panes == "G5"
        assert book["Schedule"]["B5"].value == "UNIT-2"
    finally:
        book.close()


def test_rolling_calendar_routes_past_current_and_future_records(calendar_population, tmp_path):
    template, profile = calendar_population
    book = openpyxl.load_workbook(template)
    sheet = book["Schedule"]
    sheet.insert_cols(7)
    sheet["G1"] = "Overdue"
    sheet["T1"] = ">= 2021/01"
    book.save(template)
    book.close()
    policy = profile["parameters"]["workbook_update"]
    policy["backlog_headers"] = ["Overdue"]
    policy["future_header_patterns"] = [r"^>="]
    policy["template_planning_months"].update(month_start_column=8, future_header_column=20)
    del policy["record_rules"][0]["require_date_in_planning_horizon"]
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    for month, past, present, future, expected_header in [
        ("2026/10", "2026-09-01", "2026-10-01", "2027-10-01", ">= 2027/10"),
        ("2027/01", "2026-12-01", "2027-01-01", "2028-01-01", ">= 2028/01"),
    ]:
        text = (f"Equipment: UNIT-1\nReport: {month}\n"
                f"JOB-1 2 {past}\nJOB-2 3 {present}\nJOB-3 4 {future}\n")
        output = tmp_path / f"boundaries_{month.replace('/', '_')}.xlsx"
        summary = SafeWorkbookUpdater(template, profile).execute(parse_sections(profile, text), output)
        assert summary["record_disposition_counts"] == {"INSERTED_RECORD_ROW": 3}
        assert WorkbookPreservationComparator().compare_workbooks(template, output, summary)["status"] == "PASS"
        result = openpyxl.load_workbook(output)
        try:
            sheet = result["Schedule"]
            assert sheet["T1"].value == expected_header
            assert sheet["H1"].value.strftime("%Y/%m") == month
            rows = {sheet.cell(row, 3).value: row for row in range(3, sheet.max_row + 1)}
            assert sheet.cell(rows["JOB-1"], 7).value == 2
            assert sheet.cell(rows["JOB-2"], 8).value == 3
            assert sheet.cell(rows["JOB-3"], 20).value == 4
        finally:
            result.close()
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before


def test_rolling_future_requires_a_real_matching_header(calendar_population, tmp_path):
    template, profile = calendar_population
    calendar = profile["parameters"]["workbook_update"]["template_planning_months"]
    calendar["future_header_column"] = 19
    profile["parameters"]["workbook_update"]["future_header_patterns"] = [r"^>="]
    with pytest.raises(ValueError, match="exactly one configured future header"):
        SafeWorkbookUpdater(template, profile).execute(
            parse_sections(profile, source_text("2026/10")), tmp_path / "bad.xlsx",
        )
