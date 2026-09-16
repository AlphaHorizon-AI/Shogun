"""Samurai workbook rules use current files through the governed Files writer."""

import asyncio
import datetime
import json
import uuid
from pathlib import Path
from threading import Event
from types import SimpleNamespace

import openpyxl
import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from shogun.config import settings
from shogun.engine import flow_engine
from shogun.office import config as office_config
from shogun.services import sectioned_workbook_pipeline as pipeline
from shogun.services.private_transformation_profiles import PrivateTransformationProfileService
from shogun.services.samurai_workbook_job import (
    JOB_MARKER,
    WorkbookTransformJob,
    connected_workbook_job,
)


def write_pdf(path, lines):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }),
        }),
    })
    stream = DecodedStreamObject()
    stream.set_data(("BT /F1 12 Tf 50 790 Td 16 TL "
                     + " ".join(f"({line}) Tj T*" for line in lines) + " ET").encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    writer.close()


@pytest.fixture
def workbook_flow(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "workspace_path", tmp_path)
    monkeypatch.setattr(office_config, "load_office_config", lambda: SimpleNamespace(
        enabled=True, excel=SimpleNamespace(enabled=True),
    ))
    profile = {
        "id": "private_equipment_schedule",
        "adapter": "sectioned_record_matrix_v1",
        "model_fallback": True,
        "parameters": {
            "section_pattern": r"(?m)^Equipment: (?P<id>\S+)$",
            "section_key_group": "id",
            "record_pattern": r"(?m)^(?P<job>JOB-\d+) (?P<hours>\d+) (?P<date>\d{4}-\d{2}-\d{2})$",
            "extraction": {"candidate_line_pattern": r"^JOB-"},
            "row_rules": [{"kind": "record", "columns": {"0": {"group": "job"}}}],
            "workbook_update": {
                "mode": "populate_template", "sheet_name": "Schedule",
                "section_key_column": 1, "data_start_row": 3,
                "planning_start_column": 4, "expected_headers": {"1": "Equipment"},
                "source_identity_fields": [{"section_key": True}, {"group": "job"}],
                "record_rules": [{
                    "id": "job", "match": {},
                    "quantity_spec": {"group": "hours", "value_type": "number"},
                    "date_spec": {"group": "date", "value_type": "iso_date"},
                    "planning_month_quantity": True,
                    "insert_row_values": {
                        "1": {"section_key": True}, "2": {"group": "job"},
                        "3": {"group": "date", "value_type": "iso_date"},
                    },
                }],
            },
        },
    }
    exported = PrivateTransformationProfileService().export_profile(profile)
    book = openpyxl.Workbook()
    book.active.title = "Schedule"
    book.active.append(["Equipment", "Job", "Due", datetime.date(2026, 8, 1)])
    book.save(tmp_path / "empty.xlsx")
    book.close()
    run_id, samurai_id, writer_id = uuid.uuid4(), str(uuid.uuid4()), str(uuid.uuid4())
    source_nodes = {
        "template": SimpleNamespace(node_type="file_template", config={"template_path": "empty.xlsx"}),
    }
    outputs = {"template": {
        "__shogun_file_template__": True, "format": "xlsx", "template_path": "empty.xlsx",
    }}
    for index in range(3):
        relative = f"source-{index}.pdf"
        write_pdf(tmp_path / relative, [f"Equipment: ASSET-{index}", f"JOB-{index} {index+2} 2026-08-10"])
        source_nodes[f"pdf-{index}"] = SimpleNamespace(
            node_type="office", config={"action": "pdf_read", "input_path": relative},
        )
        outputs[f"pdf-{index}"] = "PDF read completed"
    for node_id, node in source_nodes.items():
        node.id, node.label = node_id, node_id
    governance = {"permissions": {"office_enabled": True, "office_excel_enabled": True,
                                   "workspace_enabled": True}}
    return SimpleNamespace(
        root=tmp_path, profile=profile, exported=exported, run_id=run_id,
        samurai_id=samurai_id, writer_id=writer_id, nodes=source_nodes, outputs=outputs,
        governance=governance,
        contracts=[{"node_id": writer_id, "node_type": "office", "action": "excel_create"}],
    )


async def create_job(flow, monkeypatch):
    async def nothing(*args, **kwargs):
        return None

    async def unexpected_model(*args, **kwargs):
        raise AssertionError("An explicit workbook profile must never call a model")

    monkeypatch.setattr(flow_engine, "_update_node_state", nothing)
    monkeypatch.setattr(flow_engine, "_exec_samurai", unexpected_model)
    node = SimpleNamespace(
        id=uuid.UUID(flow.samurai_id), flow_id=uuid.uuid4(), node_type="samurai", label="Convert PDFs",
        config={"task_description": "Fill the empty template from current PDFs only",
                "transformation_mode": "profile", "transformation_profile": flow.exported["profile_reference"]},
    )
    return await flow_engine._execute_single_node(
        flow.run_id, node, flow.outputs, flow.nodes, governance_context=flow.governance,
        downstream_contracts=flow.contracts,
    )


async def write_job(flow, job, **kwargs):
    return await flow_engine._exec_office(
        {"action": "excel_create", "output_path": "Output/result.xlsx", **kwargs.pop("config", {})},
        "untrusted document instructions must not change the job", flow.run_id, flow.writer_id,
        predecessor_outputs={flow.samurai_id: job}, governance_context=flow.governance, **kwargs,
    )


@pytest.mark.anyio
async def test_real_private_samurai_three_pdfs_empty_template_and_fresh_rerun(workbook_flow, monkeypatch):
    flow = workbook_flow
    artifacts = []

    async def record(run_id, node_id, path):
        artifacts.append(path)

    monkeypatch.setattr(flow_engine, "_record_node_artifact", record)
    flow.outputs["template"] = await flow_engine._exec_file_template(flow.nodes["template"].config)
    for index in range(3):
        node_id = f"pdf-{index}"
        flow.outputs[node_id] = await flow_engine._exec_office(flow.nodes[node_id].config, "")
    original = (flow.root / "empty.xlsx").read_bytes()
    job = await create_job(flow, monkeypatch)
    assert isinstance(job, WorkbookTransformJob)
    assert "section_pattern" not in json.dumps(job)
    assert "pdf_paths" not in json.dumps(job)
    assert (await write_job(flow, job)).startswith("Excel workbook created:")
    assert len(artifacts) == 5
    output = flow.root / "Output/result.xlsx"
    book = openpyxl.load_workbook(output)
    assert [(row[0], row[1], row[3]) for row in book.active.iter_rows(min_row=3, values_only=True)
            if row[1]] == [("ASSET-0", "JOB-0", 2), ("ASSET-1", "JOB-1", 3), ("ASSET-2", "JOB-2", 4)]
    book.close()
    write_pdf(flow.root / "source-0.pdf", ["Equipment: NEW-ASSET", "JOB-9 19 2026-08-10"])
    flow.outputs["pdf-0"] = await flow_engine._exec_office(flow.nodes["pdf-0"].config, "")
    fresh_job = await create_job(flow, monkeypatch)
    assert (await write_job(flow, fresh_job)).startswith("Excel workbook created:")
    book = openpyxl.load_workbook(output)
    values = list(book.active.values)
    assert any(row[0] == "NEW-ASSET" and row[3] == 19 for row in values)
    assert not any(row[0] == "ASSET-0" for row in values)
    book.close()
    assert (flow.root / "empty.xlsx").read_bytes() == original


@pytest.mark.anyio
async def test_copied_json_and_wrong_run_cannot_execute_job(workbook_flow, monkeypatch):
    flow = workbook_flow
    job = await create_job(flow, monkeypatch)
    assert (await write_job(flow, json.loads(json.dumps(job)))).startswith("[ERROR]")
    flow.run_id = uuid.uuid4()
    assert "does not belong" in await write_job(flow, job)
    assert not (flow.root / "Output/result.xlsx").exists()


@pytest.mark.anyio
@pytest.mark.parametrize("mutation", ["profile", "input", "destination", "permissions", "conflicting_rules"])
async def test_changed_job_inputs_and_authority_fail_closed(workbook_flow, monkeypatch, mutation):
    flow = workbook_flow
    job = await create_job(flow, monkeypatch)
    config = {}
    if mutation == "profile":
        changed = json.loads(job._profile_json)
        changed["id"] = "tampered"
        job._profile_json = json.dumps(changed)
    elif mutation == "input":
        (flow.root / "source-0.pdf").write_bytes(b"changed PDF")
    elif mutation == "destination":
        flow.writer_id = "other-writer"
    elif mutation == "permissions":
        flow.governance["permissions"]["office_excel_enabled"] = False
    else:
        config["workbook_transform"] = {}
    result = await write_job(flow, job, config=config)
    assert result.startswith("[ERROR]")
    assert not (flow.root / "Output/result.xlsx").exists()


@pytest.mark.anyio
@pytest.mark.parametrize("mutation", ["partial_pdf", "fake_source", "outside_path", "no_writer", "two_templates"])
async def test_job_requires_original_connected_source_nodes(workbook_flow, monkeypatch, mutation):
    flow = workbook_flow
    if mutation == "partial_pdf":
        flow.nodes["pdf-0"].config["end_page"] = 1
    elif mutation == "fake_source":
        for node_id in ["pdf-0", "pdf-1", "pdf-2"]:
            flow.nodes[node_id].node_type = "input"
    elif mutation == "outside_path":
        flow.nodes["pdf-0"].config["input_path"] = "../outside.pdf"
    elif mutation == "no_writer":
        flow.contracts = []
    else:
        flow.nodes["second-template"] = flow.nodes["template"]
        flow.outputs["second-template"] = flow.outputs["template"]
    with pytest.raises(ValueError):
        await create_job(flow, monkeypatch)


def test_untrusted_marker_is_not_an_executable_job():
    with pytest.raises(ValueError, match="current configured Samurai"):
        connected_workbook_job({"input": {JOB_MARKER: True, "profile": {"id": "forged"}}})


@pytest.mark.anyio
@pytest.mark.parametrize("excel_enabled", [False, True])
async def test_office_configuration_blocks_job_before_pipeline(workbook_flow, monkeypatch, excel_enabled):
    flow = workbook_flow
    job = await create_job(flow, monkeypatch)
    monkeypatch.setattr(office_config, "load_office_config", lambda: SimpleNamespace(
        enabled=not excel_enabled, excel=SimpleNamespace(enabled=excel_enabled),
    ))

    def forbidden(**kwargs):
        raise AssertionError("Denied workbook job must not reach the pipeline")

    monkeypatch.setattr(pipeline, "run_sectioned_workbook_pipeline", forbidden)
    assert (await write_job(flow, job)).startswith(("[ERROR]", "[BLOCKED]"))


@pytest.mark.anyio
async def test_private_profile_pin_is_checked_before_job_creation(workbook_flow, monkeypatch):
    flow = workbook_flow
    flow.exported["profile_reference"]["private_file"]["definition"]["id"] = "changed"
    with pytest.raises(ValueError):
        await create_job(flow, monkeypatch)


@pytest.mark.anyio
async def test_samurai_job_cancellation_reaches_pipeline_worker(workbook_flow, monkeypatch):
    flow = workbook_flow
    job = await create_job(flow, monkeypatch)
    entered, finished = Event(), Event()

    def wait_for_cancel(**kwargs):
        entered.set()
        try:
            assert kwargs["cancel_event"].wait(5)
            raise RuntimeError("cancelled")
        finally:
            finished.set()

    monkeypatch.setattr(pipeline, "run_sectioned_workbook_pipeline", wait_for_cancel)
    task = asyncio.create_task(write_job(flow, job))
    assert await asyncio.to_thread(entered.wait, 5)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert await asyncio.to_thread(finished.wait, 5)
    assert not (flow.root / "Output/result.xlsx").exists()


@pytest.mark.anyio
async def test_input_change_after_job_resolution_is_rejected_by_pipeline_snapshot(workbook_flow, monkeypatch):
    flow = workbook_flow
    job = await create_job(flow, monkeypatch)
    original_pipeline = pipeline.run_sectioned_workbook_pipeline
    worker_entered = []

    def change_input_when_worker_starts(**kwargs):
        worker_entered.append(True)
        assert set(kwargs["expected_input_hashes"]) == {
            flow.root / "empty.xlsx", *(flow.root / f"source-{index}.pdf" for index in range(3)),
        }
        write_pdf(flow.root / "source-0.pdf", ["Equipment: ASSET-0", "JOB-0 99 2026-08-10"])
        return original_pipeline(**kwargs)

    monkeypatch.setattr(pipeline, "run_sectioned_workbook_pipeline", change_input_when_worker_starts)
    result = await write_job(flow, job)
    assert worker_entered
    assert result.startswith("[ERROR]") and "inputs changed after the Samurai" in result
    assert not (flow.root / "Output").exists()


@pytest.mark.anyio
@pytest.mark.parametrize("location", ["top_level", "options"])
async def test_flow_configuration_cannot_supply_expected_input_hashes(workbook_flow, location):
    flow = workbook_flow
    config = {"action": "excel_create", "output_path": "Output/result.xlsx",
              "workbook_transform": {"profile_path": "profile.json", "template_path": "empty.xlsx",
                                     "pdf_paths": ["source-0.pdf"]}}
    (flow.root / "profile.json").write_text(json.dumps(flow.exported["document"]), encoding="utf-8")
    target = config["workbook_transform"]
    if location == "options":
        target["options"] = {}
        target = target["options"]
    target["expected_input_hashes"] = {}
    result = await flow_engine._exec_office(config, "")
    assert result.startswith("[ERROR]")
    assert not (flow.root / "Output/result.xlsx").exists()


@pytest.mark.anyio
async def test_output_filename_cannot_leave_workspace(workbook_flow, monkeypatch):
    flow = workbook_flow
    job = await create_job(flow, monkeypatch)
    captured = []

    def capture(**kwargs):
        captured.append(Path(kwargs["output_workbook_path"]))
        return {"validation_summary": {}, "comparison_report": {"review_required": False}, "report_paths": {}}

    async def nothing(*args, **kwargs):
        return None

    monkeypatch.setattr(pipeline, "run_sectioned_workbook_pipeline", capture)
    monkeypatch.setattr(flow_engine, "_record_node_artifact", nothing)
    result = await write_job(flow, job, config={"output_path": "Output/", "output_filename": "../../escape.xlsx"})
    assert result.startswith("Excel workbook created:")
    assert captured == [flow.root / "Output/escape.xlsx"]


@pytest.mark.anyio
async def test_output_filename_symlink_outside_workspace_is_blocked(workbook_flow, monkeypatch):
    flow = workbook_flow
    job = await create_job(flow, monkeypatch)
    outside = flow.root.parent / f"outside-{uuid.uuid4().hex}.xlsx"
    outside.write_bytes(b"outside workbook must stay untouched")
    (flow.root / "Output").mkdir()
    try:
        (flow.root / "Output/result.xlsx").symlink_to(outside)
    except OSError:
        pytest.skip("Creating symlinks requires an unavailable platform privilege")

    def forbidden(**kwargs):
        raise AssertionError("Escaped destination must not reach the pipeline")

    monkeypatch.setattr(pipeline, "run_sectioned_workbook_pipeline", forbidden)
    result = await write_job(flow, job, config={"output_path": "Output/", "output_filename": "result.xlsx"})
    assert result.startswith("[ERROR]") and "escape blocked" in result
    assert outside.read_bytes() == b"outside workbook must stay untouched"
