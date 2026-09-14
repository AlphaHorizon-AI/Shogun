"""Independent regression tests using an unrelated text-only maintenance workflow."""

from copy import deepcopy
from datetime import datetime

import openpyxl
import pytest

from shogun.services.private_transformation_profiles import PrivateTransformationProfileService
from shogun.services.sectioned_workbook_pipeline import SectionedLayoutParser, run_sectioned_workbook_pipeline
from shogun.services.sectioned_workbook_updater import SafeWorkbookUpdater, WorkbookPreservationComparator
from shogun.services.transformation_profile_registry import profile_content_hash


@pytest.fixture
def equipment_profile():
    return {
        "id": "private_maintenance_notes_v1",
        "adapter": "sectioned_record_matrix_v1",
        "model_fallback": False,
        "parameters": {
            "section_pattern": r"(?m)^Asset (?P<section_id>\S+)$",
            "record_pattern": r"^Task (?P<reference>\S+) (?P<note>.+)$",
            "extraction": {"candidate_line_pattern": r"^Task\b"},
            "row_rules": [{"kind": "record", "columns": {"0": {"group": "note"}}}],
            "workbook_update": {
                "header_row": 3,
                "data_start_row": 5,
                "section_key_column": 4,
                "require_planning_months": False,
                "expected_headers": {"4": "Asset", "2": "Task", "6": "Work note"},
                "source_identity_fields": [{"section_key": True}, {"group": "reference"}],
                "record_rules": [
                    {
                        "id": "maintenance_note",
                        "reference_spec": {"group": "reference"},
                        "reference_column": 2,
                        "insert_row_values": {
                            "2": {"group": "reference"},
                            "4": {"section_key": True},
                            "6": {"group": "note"},
                        },
                    }
                ],
            },
        },
    }


def template(path, reference=None):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Maintenance"
    sheet["A1"] = "Supervisor worksheet"
    sheet.append([])
    for col, value in {2: "Task", 4: "Asset", 6: "Work note"}.items():
        sheet.cell(3, col).value = value
    sheet["D5"] = "PUMP-X"
    sheet["B5"] = reference
    sheet["F5"] = "Preserve supervisor's note"
    sheet.freeze_panes = "C5"
    book.save(path)
    book.close()
    return path


def values(path):
    book = openpyxl.load_workbook(path)
    try:
        return list(book["Maintenance"].values)
    finally:
        book.close()


def test_text_workflow_has_no_quantity_date_or_months_and_reruns(equipment_profile, tmp_path):
    service = PrivateTransformationProfileService()
    profile = service.import_document(service.export_profile(equipment_profile)["document"])["document"]["profile"]
    sections = SectionedLayoutParser(profile).parse_pages([(1, "Asset PUMP-X\nTask N-1 Inspect gasket")], "notes.pdf")
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    summary = SafeWorkbookUpdater(template(tmp_path / "base.xlsx"), profile).execute(sections, first, "Maintenance")
    assert summary["normalized_record_count"] == summary["accounted_record_count"] == 1
    assert summary["total_rows_inserted"] == 1
    assert values(first)[4][5] == "Preserve supervisor's note"
    assert values(first)[5][1:6] == ("N-1", None, "PUMP-X", None, "Inspect gasket")
    rerun = SafeWorkbookUpdater(first, profile).execute(sections, second, "Maintenance")
    assert rerun["total_rows_inserted"] == 0 and not rerun["allowed_changes"]
    assert values(first) == values(second)


def test_identifiers_are_exact_unless_profile_explicitly_requests_normalization(equipment_profile, tmp_path):
    sections = SectionedLayoutParser(equipment_profile).parse_pages(
        [(1, "Asset PUMP-X\nTask 000123 Distinct work ticket")],
        "notes.pdf",
    )
    first = tmp_path / "first.xlsx"
    summary = SafeWorkbookUpdater(template(tmp_path / "base.xlsx", "123"), equipment_profile).execute(
        sections,
        first,
        "Maintenance",
    )
    assert summary["total_rows_inserted"] == 1
    assert [row[1] for row in values(first)[4:]] == ["123", "000123"]


@pytest.mark.parametrize("column", [True, 0, -1, 16385, 2.5])
def test_invalid_columns_rejected_at_private_import(equipment_profile, column):
    profile = deepcopy(equipment_profile)
    profile["parameters"]["workbook_update"]["record_rules"][0]["reference_column"] = column
    with pytest.raises(ValueError):
        PrivateTransformationProfileService().export_profile(profile)


def test_unknown_nested_policy_is_rejected(equipment_profile):
    equipment_profile["parameters"]["workbook_update"]["record_rules"][0]["overwrite_everything"] = True
    with pytest.raises(ValueError, match="Unsupported|unsupported"):
        PrivateTransformationProfileService().export_profile(equipment_profile)


@pytest.mark.parametrize("tamper", ["hash", "extra_column", "empty_version", "nan_value", "changed_profile"])
def test_provenance_rejects_corruption_and_profile_changes(equipment_profile, tmp_path, tamper):
    sections = SectionedLayoutParser(equipment_profile).parse_pages(
        [(1, "Asset PUMP-X\nTask N-1 Inspect gasket")],
        "notes.pdf",
    )
    first = tmp_path / "first.xlsx"
    SafeWorkbookUpdater(template(tmp_path / "base.xlsx"), equipment_profile).execute(sections, first)
    book = openpyxl.load_workbook(first)
    meta = book["_shogun_workflow_provenance"]
    assert meta["F2"].value == profile_content_hash(equipment_profile)
    if tamper == "hash":
        meta["F2"] = "0" * 64
    elif tamper == "extra_column":
        meta["G1"] = "unexpected"
    elif tamper == "empty_version":
        meta["A2"] = None
    elif tamper == "nan_value":
        meta["E2"] = '{"6": NaN}'
    else:
        equipment_profile["parameters"]["workbook_update"]["record_rules"][0]["insert_row_values"]["6"] = {
            "literal": "Different rule",
        }
    book.save(first)
    book.close()
    output = tmp_path / "rejected.xlsx"
    with pytest.raises(ValueError, match="[Pp]rovenance|profile"):
        SafeWorkbookUpdater(first, equipment_profile).execute(sections, output)
    assert not output.exists()


def test_allowed_number_format_does_not_hide_font_changes(equipment_profile, tmp_path):
    baseline = template(tmp_path / "base.xlsx", "N-1")
    book = openpyxl.load_workbook(baseline)
    book["Maintenance"]["C5"] = datetime(2026, 9, 14)
    book.save(baseline)
    book.close()
    sections = SectionedLayoutParser(equipment_profile).parse_pages(
        [(1, "Asset PUMP-X\nTask N-1 Inspect gasket")],
        "notes.pdf",
    )
    output = tmp_path / "generated.xlsx"
    summary = SafeWorkbookUpdater(baseline, equipment_profile).execute(sections, output)
    # Add a declared number-format allowance, then change an unrelated style property.
    summary["allowed_style_changes"] = [
        {
            "original_row": 5,
            "column": 3,
            "property": "number_format",
            "new_value": "yyyy-mm-dd",
        }
    ]
    book = openpyxl.load_workbook(output)
    book["Maintenance"]["C5"].number_format = "yyyy-mm-dd"
    book["Maintenance"]["C5"].font = openpyxl.styles.Font(bold=True, color="FF0000")
    book.save(output)
    book.close()
    comparison = WorkbookPreservationComparator().compare_workbooks(baseline, output, summary)
    assert comparison["status"] == "FAIL"
    assert any(issue["type"] == "cell_style" for issue in comparison["issues"])


def test_record_rules_respect_section_conditions(equipment_profile, tmp_path):
    rules = equipment_profile["parameters"]["workbook_update"]["record_rules"]
    second = deepcopy(rules[0])
    second["id"] = "other_asset"
    rules.append(second)
    for rule, key in zip(rules, ["PUMP-X", "PUMP-Y"]):
        rule["when"] = {"field": "section_id", "operator": "equals", "value": key}
    sections = SectionedLayoutParser(equipment_profile).parse_pages(
        [(1, "Asset PUMP-X\nTask N-1 Inspect gasket")],
        "notes.pdf",
    )
    assert not sections[0].records[0].is_invalid
    output = tmp_path / "generated.xlsx"
    summary = SafeWorkbookUpdater(template(tmp_path / "base.xlsx"), equipment_profile).execute(sections, output)
    assert summary["total_rows_inserted"] == 1


@pytest.mark.parametrize("target", ["section_selection", "rule"])
def test_misspelled_condition_does_not_silently_select_records(equipment_profile, target):
    condition = {"field": "section_id", "equals": "PUMP-X"}
    params = equipment_profile["parameters"]
    if target == "section_selection":
        params[target] = condition
    else:
        params["workbook_update"]["record_rules"][0]["when"] = condition
    with pytest.raises(ValueError, match="[Uu]nsupported"):
        PrivateTransformationProfileService().export_profile(equipment_profile)


@pytest.mark.parametrize("field", ["quantity", "alternative_quantity", "date"])
def test_writer_revalidates_changed_raw_values(equipment_profile, tmp_path, field):
    rule = equipment_profile["parameters"]["workbook_update"]["record_rules"][0]
    if field == "date":
        rule["date_spec"] = {"group": "note", "value_type": "iso_date"}
        source = "Task D-1 2026-09-14"
    else:
        rule["quantity_spec"] = {"group": "note", "value_type": "number"}
        source = "Task 3 3"
        if field == "alternative_quantity":
            rule["alternative_quantity_spec"] = {"group": "reference", "value_type": "number"}
    sections = SectionedLayoutParser(equipment_profile).parse_pages([(1, "Asset PUMP-X\n" + source)], "notes.pdf")
    record = sections[0].records[0]
    assert not record.is_invalid
    record.raw["reference" if field == "alternative_quantity" else "note"] = "malformed"
    result = SafeWorkbookUpdater(template(tmp_path / "base.xlsx"), equipment_profile).execute(
        sections,
        tmp_path / "reviewed.xlsx",
    )
    assert result["total_rows_inserted"] == 0
    assert any(event.get("record_id") and event["confidence"] == "REVIEW" for event in result["events"])


def test_profile_input_cannot_be_overwritten_by_a_report(equipment_profile, tmp_path):
    protected = tmp_path / "result_normalized_pdf_materials.json"
    protected.write_text("private profile source", encoding="utf-8")
    pdf = tmp_path / "input.pdf"
    pdf.write_bytes(b"%PDF-fixture")
    with pytest.raises(ValueError, match="aliases input"):
        run_sectioned_workbook_pipeline(
            equipment_profile,
            [pdf],
            template(tmp_path / "base.xlsx"),
            tmp_path / "result.xlsx",
            profile_source_path=protected,
        )
    assert protected.read_text(encoding="utf-8") == "private profile source"


@pytest.mark.parametrize("existing", [None, 0, 7])
def test_section_rule_preserves_values_and_regex_quantifiers(equipment_profile, tmp_path, existing):
    equipment_profile["parameters"]["workbook_update"]["section_rules"] = [
        {
            "id": "inspection_score",
            "match_column": 6,
            "value_column": 1,
            "match_pattern": r"Ticket-\d{3}-{section_key}",
            "value_spec": {"literal": 5},
        }
    ]
    baseline = template(tmp_path / "base.xlsx")
    book = openpyxl.load_workbook(baseline)
    book["Maintenance"]["F5"] = "Ticket-123-PUMP-X"
    book["Maintenance"]["A5"] = existing
    book.save(baseline)
    book.close()
    sections = SectionedLayoutParser(equipment_profile).parse_pages([(1, "Asset PUMP-X")], "notes.pdf")
    output = tmp_path / "result.xlsx"
    result = SafeWorkbookUpdater(baseline, equipment_profile).execute(sections, output)
    assert result["total_rows_inserted"] == 0
    assert values(output)[4][0] == (5 if existing is None else existing)
