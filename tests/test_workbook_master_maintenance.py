"""Planner details survive current-source refreshes of an owned master workbook."""

import copy
import datetime
import hashlib
import json

import openpyxl
import pytest
from openpyxl.comments import Comment

from shogun.services.sectioned_workbook_pipeline import SectionedLayoutParser
from shogun.services.sectioned_workbook_updater import SafeWorkbookUpdater, WorkbookPreservationComparator
from shogun.services.workbook_master_maintenance import HISTORY_SHEET, MISSING_COMMENT, STATE_SHEET


@pytest.fixture
def master_setup(tmp_path):
    seed = tmp_path / "seed.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Plan"
    sheet.append(["Description", "Article", "Family", "Planner", "Reference", "Quantity",
                  "Customer", "Notes", "Stage", "Due"])
    for column, month in enumerate((8, 9, 10), 11):
        sheet.cell(1, column, datetime.datetime(2026, month, 1)).number_format = "mmm yyyy"
    sheet.cell(3, 2).number_format = "@"
    sheet.column_dimensions["G"].width = 22
    sheet.freeze_panes = "A3"
    book.save(seed)
    book.close()
    profile = {
        "id": "maintained_plan", "adapter": "sectioned_record_matrix_v1",
        "parameters": {
            "section_pattern": (
                r"(?m)^Part: (?P<id>\S+) Group: (?P<family_group>\d+) "
                r"Size: (?P<size>\d+) Family: (?P<family>\S+)$"
            ),
            "section_key_group": "id",
            "section_fields": [
                {"target": "month", "pattern": r"Report: (?P<value>\d{4}/\d{2})", "group": "value"},
                {"target": "stock", "pattern": r"Stock: (?P<value>\d+)",
                 "group": "value", "value_type": "number"},
            ],
            "record_pattern": (
                r"(?m)^(?P<kind>JOB|DEM) (?P<reference>\S+) (?P<quantity>\d+) "
                r"(?P<month>\d{4}/\d{2}) (?P<date>\d{4}-\d{2}-\d{2})$"
            ),
            "extraction": {"candidate_line_pattern": r"^(?:JOB|DEM) "},
            "row_rules": [{"kind": "record", "columns": {"0": {"group": "reference"}}}],
            "workbook_update": {
                "mode": "maintain_master", "sheet_name": "Plan", "section_key_column": 2,
                "data_start_row": 3, "expected_headers": {"2": "Article"}, "planning_start_column": 11,
                "require_planning_months": True, "template_section_row_policy": "merge_first_detail",
                "template_planning_months": {"start_month_spec": {"field": "month"}, "count": 3},
                "template_section_values": {"1": {"section_key": True}, "3": {"field": "family"}},
                "template_section_sort": [
                    {"field": "family_group", "value_type": "number"},
                    {"field": "size", "value_type": "number"}, {"field": "family"}, {"section_key": True},
                ],
                "source_identity_fields": [
                    {"section_key": True}, {"group": "kind"}, {"group": "reference"},
                ],
                "section_rules": [{
                    "id": "stock", "value_spec": {"field": "stock"}, "value_column": 6,
                    "match_column": 5, "match_pattern": "Stock", "zero_value_skip_insert": True,
                    "insert_row_values": {
                        "2": {"section_key": True}, "5": {"literal": "Stock"}, "6": {"field": "stock"},
                    },
                }],
                "record_rules": [{
                    "id": "order", "match": {"kind": "JOB"}, "reference_column": 5,
                    "reference_spec": {"group": "reference"},
                    "quantity_spec": {"group": "quantity", "value_type": "number"},
                    "date_spec": {"group": "date", "value_type": "iso_date"},
                    "insert_row_values": {
                        "2": {"section_key": True}, "5": {"group": "reference"},
                        "6": {"group": "quantity", "value_type": "number"},
                        "10": {"group": "date", "value_type": "iso_date"},
                    },
                }, {
                    "id": "demand", "match": {"kind": "DEM"},
                    "reference_spec": {"group": "reference"},
                    "quantity_spec": {"group": "quantity", "value_type": "number"},
                    "date_spec": {"group": "date", "value_type": "iso_date"},
                    "planning_month_spec": {"group": "month"}, "planning_month_quantity": True,
                    "require_date_in_planning_horizon": True,
                    "insert_row_values": {
                        "2": {"section_key": True}, "5": {"group": "reference"},
                        "10": {"group": "date", "value_type": "iso_date"},
                    },
                }],
            },
        },
    }
    return seed, profile


def section(key="A", report="2026/08", stock=1, size=100, records=None):
    return (
        f"Part: {key} Group: 1 Size: {size} Family: {key}\nReport: {report}\nStock: {stock}\n"
        + (records if records is not None else "JOB P1 30 2026/08 2026-08-17\nDEM D1 12 2026/10 2026-10-05\n")
    )


def run_master(source, profile, text, output):
    sections = SectionedLayoutParser(profile).parse_pages(list(enumerate(text.split("\f"), 1)), "current.pdf")
    return SafeWorkbookUpdater(source, profile).execute(sections, output)


def row_for(sheet, reference):
    return next(row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 5).value == reference)


def test_seed_and_idempotent_rerun(master_setup, tmp_path):
    seed, profile = master_setup
    before = seed.read_bytes()
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    initial = run_master(seed, profile, section(), first)
    assert initial["total_rows_inserted"] == 3
    repeated = run_master(first, profile, section(), second)
    assert repeated["total_rows_inserted"] == 0
    assert repeated["record_disposition_counts"] == {"MATCHED_SOURCE_RECORD": 2}
    assert seed.read_bytes() == before
    book = openpyxl.load_workbook(second)
    try:
        assert book["Plan"].freeze_panes == "A3"
        assert book[STATE_SHEET].sheet_state == "veryHidden"
        assert book[HISTORY_SHEET].max_row == 1
        assert book["Plan"].cell(row_for(book["Plan"], "D1"), 13).value == 12
    finally:
        book.close()


def test_august_to_october_updates_owned_values_preserves_planner_cells_and_archives_departures(master_setup, tmp_path):
    seed, profile = master_setup
    first, second = tmp_path / "august.xlsx", tmp_path / "october.xlsx"
    run_master(seed, profile, section(), first)
    book = openpyxl.load_workbook(first)
    sheet = book["Plan"]
    order, demand = row_for(sheet, "P1"), row_for(sheet, "D1")
    sheet.cell(order, 6, 99)  # Planner overrides generated order quantity.
    sheet.cell(order, 7, "Customer ACME")
    sheet.cell(order, 8, "Keep my planning note")
    sheet.cell(order, 9, "Avo 17")
    sheet.cell(order, 11, 4)  # Manual August quantity must be archived, not relabelled October.
    sheet.cell(demand, 10).comment = Comment("Planner date note", "Planner")
    sheet.cell(demand, 12, 9)  # A separate manual September cell.
    book.save(first)
    book.close()
    before = first.read_bytes()
    updated = run_master(first, profile, section(report="2026/10", stock=0, records=(
        "JOB P1 35 2026/10 2026-10-12\nDEM D1 6 2026/12 2026-12-07\n"
    )), second)
    assert updated["total_rows_inserted"] == 0
    assert updated["status_counts"]["PRESERVED_MANUAL_SOURCE_VALUE"] == 1
    book = openpyxl.load_workbook(second)
    try:
        sheet = book["Plan"]
        assert sheet.cell(order, 6).value == 99
        assert [sheet.cell(order, col).value for col in (7, 8, 9)] == [
            "Customer ACME", "Keep my planning note", "Avo 17",
        ]
        assert sheet.cell(order, 10).value == datetime.datetime(2026, 10, 12)
        assert sheet.cell(row_for(sheet, "Stock"), 6).value == 0
        assert sheet.cell(demand, 10).value == datetime.datetime(2026, 12, 7)
        assert sheet.cell(demand, 10).comment.text == "Planner date note"
        assert sheet.cell(demand, 11).value is None  # Source demand moved from October to December.
        assert sheet.cell(demand, 13).value == 6
        assert sheet.cell(order, 11).value is None
        archived = list(book[HISTORY_SHEET].iter_rows(min_row=2, values_only=True))
        assert {(row[6], json.loads(row[7])["value"]) for row in archived} == {("2026/08", 4), ("2026/09", 9)}
        assert sheet.freeze_panes == "A3"
    finally:
        book.close()
    assert first.read_bytes() == before


def test_new_material_inserts_without_reordering_existing_groups_and_missing_orders_remain_marked(
    master_setup, tmp_path,
):
    seed, profile = master_setup
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    run_master(seed, profile, section() + "\f" + section("C", size=300, records=""), first)
    book = openpyxl.load_workbook(first)
    order = row_for(book["Plan"], "P1")
    book["Plan"].cell(order, 8, "Do not delete")
    book["Plan"].cell(order, 2).comment = Comment("Planner article note", "Planner")
    book.save(first)
    book.close()
    result = run_master(first, profile, (
        section(records="DEM D1 12 2026/10 2026-10-05\n")
        + "\f" + section("B", size=200, records="")
        + "\f" + section("C", size=300, records="")
    ), second)
    assert result["status_counts"]["MISSING_SOURCE_RECORD_RETAINED"] == 1
    book = openpyxl.load_workbook(second)
    try:
        sheet = book["Plan"]
        keys = [sheet.cell(row, 2).value for row in range(3, sheet.max_row + 1)]
        assert list(dict.fromkeys(key for key in keys if key)) == ["A", "B", "C"]
        assert sheet.cell(row_for(sheet, "P1"), 8).value == "Do not delete"
        comment = sheet.cell(row_for(sheet, "P1"), 2).comment
        assert comment.text == "Planner article note\n" + MISSING_COMMENT
        assert comment.author == "Planner"
    finally:
        book.close()


def test_year_roll_moves_manual_month_value_by_month(master_setup, tmp_path):
    seed, profile = master_setup
    first, second = tmp_path / "november.xlsx", tmp_path / "december.xlsx"
    run_master(seed, profile, section(report="2026/11", records=""), first)
    book = openpyxl.load_workbook(first)
    stock = row_for(book["Plan"], "Stock")
    book["Plan"].cell(stock, 13, 42)  # January 2027 is currently the third month.
    book.save(first)
    book.close()
    run_master(first, profile, section(report="2026/12", records=""), second)
    book = openpyxl.load_workbook(second)
    try:
        sheet = book["Plan"]
        assert [sheet.cell(1, col).value.strftime("%Y/%m") for col in range(11, 14)] == [
            "2026/12", "2027/01", "2027/02",
        ]
        assert sheet.cell(stock, 12).value == 42
        assert sheet.cell(stock, 13).value is None
    finally:
        book.close()


def test_blank_departing_month_comments_are_archived_and_absent_material_remains_visible(master_setup, tmp_path):
    seed, profile = master_setup
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    run_master(seed, profile, section(stock=0, records="") + "\f" + section("B", records=""), first)
    book = openpyxl.load_workbook(first)
    article = next(row for row in range(3, book["Plan"].max_row + 1) if book["Plan"].cell(row, 2).value == "A")
    book["Plan"].cell(article, 11).comment = Comment("August note without quantity", "Planner")
    book.save(first)
    book.close()
    summary = run_master(first, profile, section("B", report="2026/10", records=""), second)
    assert summary["status_counts"]["MISSING_SOURCE_SECTION_RETAINED"] == 1
    book = openpyxl.load_workbook(second)
    try:
        assert book["Plan"].cell(article, 2).value == "A"
        assert MISSING_COMMENT in book["Plan"].cell(article, 2).comment.text
        archived = list(book[HISTORY_SHEET].iter_rows(min_row=2, values_only=True))
        payload = json.loads(next(row[7] for row in archived if row[6] == "2026/08"))
        assert payload["value"] is None
        assert payload["comment"] == {"text": "August note without quantity", "author": "Planner"}
    finally:
        book.close()


def test_master_profile_change_is_rejected_without_touching_files(master_setup, tmp_path):
    seed, profile = master_setup
    master, output = tmp_path / "master.xlsx", tmp_path / "rejected.xlsx"
    run_master(seed, profile, section(), master)
    before = hashlib.sha256(master.read_bytes()).hexdigest()
    changed = copy.deepcopy(profile)
    changed["parameters"]["workbook_update"]["template_section_values"]["1"] = {"literal": "Changed"}
    with pytest.raises(ValueError, match="provenance metadata"):
        run_master(master, changed, section(), output)
    assert not output.exists()
    assert hashlib.sha256(master.read_bytes()).hexdigest() == before


@pytest.mark.parametrize("change", ["insert_row", "order_reference"])
def test_manual_row_structure_and_reference_changes_fail_safely(master_setup, tmp_path, change):
    seed, profile = master_setup
    master, output = tmp_path / "master.xlsx", tmp_path / "rejected.xlsx"
    run_master(seed, profile, section(), master)
    book = openpyxl.load_workbook(master)
    sheet = book["Plan"]
    if change == "insert_row":
        sheet.insert_rows(4)
        sheet["B4"] = "A"
    else:
        sheet.cell(row_for(sheet, "P1"), 5, "Changed reference")
    book.save(master)
    book.close()
    before = master.read_bytes()
    with pytest.raises(ValueError, match="ownership|positions"):
        run_master(master, profile, section(), output)
    assert master.read_bytes() == before
    assert not output.exists()


def test_hyperlinks_follow_month_and_row_insertion(master_setup, tmp_path):
    seed, profile = master_setup
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    run_master(seed, profile, section(report="2026/11", records=""), first)
    book = openpyxl.load_workbook(first)
    sheet = book["Plan"]
    stock = row_for(sheet, "Stock")
    sheet.cell(stock, 13, "January details").hyperlink = "https://example.com/january"
    sheet.cell(stock, 7, "Customer details").hyperlink = "https://example.com/customer"
    book.save(first)
    book.close()
    run_master(first, profile, (
        section(report="2026/12", records="") + "\f"
        + section("B", report="2026/12", size=90, records="")
    ), second)
    book = openpyxl.load_workbook(second)
    try:
        sheet = book["Plan"]
        row = next(row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 2).value == "A")
        assert row == stock + 1
        assert sheet.cell(row, 12).value == "January details"
        assert sheet.cell(row, 12).hyperlink.target == "https://example.com/january"
        assert sheet.cell(row, 12).hyperlink.ref == f"L{row}"
        assert sheet.cell(row, 7).hyperlink.target == "https://example.com/customer"
        assert sheet.cell(row, 7).hyperlink.ref == f"G{row}"
    finally:
        book.close()


@pytest.mark.parametrize("initial_component", ["A", ""])
def test_zero_stock_optional_component_blanks_round_trip_when_new_material_shifts_rows(
    master_setup, tmp_path, initial_component,
):
    seed, profile = master_setup
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    original_sections = SectionedLayoutParser(profile).parse_pages([(1, section(records=""))], "august.pdf")
    original_sections[0].fields["family"] = initial_component
    SafeWorkbookUpdater(seed, profile).execute(original_sections, first)
    current_sections = SectionedLayoutParser(profile).parse_pages([
        (1, section(stock=0, records="")), (2, section("B", size=90, records="")),
    ], "october.pdf")
    current_sections[0].fields["family"] = ""
    summary = SafeWorkbookUpdater(first, profile).execute(current_sections, second)
    comparison = WorkbookPreservationComparator().compare_workbooks(first, second, summary)
    assert comparison["status"] == "PASS"
    assert summary["total_rows_inserted"] == 1
    assert all(change["new_value"] != "" for change in summary["allowed_changes"])
    book = openpyxl.load_workbook(second)
    try:
        sheet = book["Plan"]
        old_article_row = next(row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 2).value == "A")
        assert old_article_row == 4
        assert sheet.cell(old_article_row, 3).value is None
        assert sheet.cell(old_article_row, 6).value == 0
    finally:
        book.close()


def test_failed_preservation_check_does_not_publish_or_modify_master(master_setup, tmp_path, monkeypatch):
    seed, profile = master_setup
    master, output = tmp_path / "master.xlsx", tmp_path / "existing-output.xlsx"
    run_master(seed, profile, section(), master)
    output.write_bytes(b"existing output remains intact")
    before = master.read_bytes()
    compare = WorkbookPreservationComparator.compare_workbooks

    def reject_master(self, baseline, generated, summary):
        if summary.get("mode") == "maintain_master":
            return {"status": "FAIL", "issues": [{"type": "test_failure"}]}
        return compare(self, baseline, generated, summary)

    monkeypatch.setattr(WorkbookPreservationComparator, "compare_workbooks", reject_master)
    with pytest.raises(ValueError, match="Master preservation verification failed"):
        run_master(master, profile, section(), output)
    assert master.read_bytes() == before
    assert output.read_bytes() == b"existing output remains intact"
