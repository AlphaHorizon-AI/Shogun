"""Empty templates are populated from current sources, without material fixtures."""

import copy
import datetime
import hashlib
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from shogun.services.sectioned_workbook_pipeline import SectionedLayoutParser
from shogun.services.sectioned_workbook_updater import (
    SafeWorkbookUpdater,
    WorkbookPreservationComparator,
    validate_workbook_update_profile,
)


@pytest.fixture
def population(tmp_path):
    template = tmp_path / "empty.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Schedule"
    sheet.append([
        "Description", "Equipment", "Job", "Hours", "Due date", "Overdue",
        datetime.date(2026, 7, 1), datetime.date(2026, 8, 1), ">= Sep 2026",
    ])
    sheet.cell(3, 2).fill = PatternFill("solid", fgColor="AADDCC")
    sheet.cell(12, 2).fill = PatternFill("solid", fgColor="AADDCC")
    sheet.row_dimensions[3].height = 23
    book.save(template)
    book.close()
    profile = {
        "id": "equipment_schedule",
        "adapter": "sectioned_record_matrix_v1",
        "parameters": {
            "section_pattern": r"(?m)^Equipment: (?P<id>\S+) Size: (?P<size>\d+)$",
            "section_key_group": "id",
            "record_pattern": r"(?m)^(?P<job>JOB-\d+) (?P<hours>\d+) (?P<date>\d{4}-\d{2}-\d{2})$",
            "extraction": {"candidate_line_pattern": r"^JOB-"},
            "row_rules": [{"kind": "record", "columns": {"0": {"group": "job"}}}],
            "workbook_update": {
                "mode": "populate_template",
                "sheet_name": "Schedule",
                "section_key_column": 2,
                "data_start_row": 3,
                "expected_headers": {"2": "Equipment"},
                "planning_start_column": 6,
                "backlog_headers": ["Overdue"],
                "future_header_patterns": [r"^>="],
                "require_planning_months": True,
                "template_section_values": {"1": {"field": "size"}},
                "template_section_sort": [
                    {"field": "size", "value_type": "number"}, {"section_key": True},
                ],
                "source_identity_fields": [{"section_key": True}, {"group": "job"}],
                "record_rules": [{
                    "id": "job",
                    "match": {},
                    "quantity_spec": {"group": "hours", "value_type": "number"},
                    "date_spec": {"group": "date", "value_type": "iso_date"},
                    "planning_month_quantity": True,
                    "insert_row_values": {
                        "2": {"section_key": True},
                        "3": {"group": "job"},
                        "5": {"group": "date", "value_type": "iso_date"},
                    },
                }],
            },
        },
    }
    return template, profile


def execute(population, tmp_path, text, filename="result.xlsx"):
    template, profile = population
    pages = list(enumerate(text.split("\f"), 1))
    sections = SectionedLayoutParser(profile).parse_pages(pages, source_file="current.pdf")
    path = tmp_path / filename
    summary = SafeWorkbookUpdater(template, profile).execute(sections, path)
    return path, summary


def schedule_rows(path):
    book = openpyxl.load_workbook(path)
    try:
        return [
            row for row in book["Schedule"].iter_rows(min_row=3, max_col=9, values_only=True)
            if any(value is not None for value in row)
        ]
    finally:
        book.close()


@pytest.fixture
def dependency_population(population):
    _, profile = population
    profile["parameters"]["section_pattern"] = (
        r"(?m)^Equipment: (?P<id>\S+) Size: (?P<size>\d+)"
        r"(?: Dependency: (?P<dependency>\S+))?$"
    )
    profile["parameters"]["section_order"] = {
        "dependencies_before": {"fields": ["dependency"], "missing": "error"},
    }
    return population


@pytest.fixture
def component_population(population):
    _, profile = population
    profile["parameters"]["selector_fields"] = [{
        "target": "component",
        "scope_pattern": r"(?s)\A(?P<body>.*)\Z",
        "line_pattern": r"(?m)^Component: (?P<value>\S+) (?P<text>.+)$",
        "minimum_matches": 1, "maximum_matches": 1,
        "distinct": True, "on_cardinality_mismatch": "preserve",
    }]
    profile["parameters"]["resolution_groups"] = [{
        "name": "equipment_component", "targets": ["component"],
        "status_target": "component_status", "requires_review_target": "component_review",
    }]
    return population


@pytest.mark.parametrize("components,use_resolution,review_required", [
    ("", True, True),
    ("Component: PART-A lower\nComponent: PART-B lower\n", True, True),
    ("Component: PART-A lower\nComponent: PART-B lower\n", False, True),
    ("Component: PART-A lower\n", True, False),
])
def test_selected_source_field_uncertainty_requires_review_without_losing_records(
    component_population, tmp_path, components, use_resolution, review_required,
):
    template, profile = component_population
    if not use_resolution:
        del profile["parameters"]["resolution_groups"]
    before = template.read_bytes()
    path, summary = execute(component_population, tmp_path, (
        f"Equipment: UNIT-100 Size: 100\n{components}JOB-1 7 2026-07-06\n"
    ))
    comparison = WorkbookPreservationComparator().compare_workbooks(template, path, summary)
    assert comparison["status"] == "PASS"
    assert comparison["review_required"] is review_required
    assert any(event["status"] == "SOURCE_FIELD_REVIEW" for event in summary["events"]) is review_required
    assert summary["normalized_record_count"] == summary["accounted_record_count"] == 1
    assert schedule_rows(path) == [
        ("100", "UNIT-100", None, None, None, None, None, None, None),
        (None, "UNIT-100", "JOB-1", None, datetime.datetime(2026, 7, 6), None, 7, None, None),
    ]
    assert template.read_bytes() == before


@pytest.mark.parametrize("components,use_resolution", [
    ("", True),
    ("Component: PART-A lower\nComponent: PART-B lower\n", True),
    ("Component: PART-A lower\nComponent: PART-B lower\n", False),
])
def test_excluded_source_field_uncertainty_does_not_flag_selected_output(
    component_population, tmp_path, components, use_resolution,
):
    template, profile = component_population
    if not use_resolution:
        del profile["parameters"]["resolution_groups"]
    profile["parameters"]["section_selection"] = {"field": "id", "operator": "equals", "value": "KEEP-100"}
    path, summary = execute(component_population, tmp_path, (
        "Equipment: KEEP-100 Size: 100\nComponent: PART-A lower\nJOB-1 7 2026-07-06\n\f"
        f"Equipment: EXCLUDED-200 Size: 200\n{components}JOB-2 9 2026-07-13\n"
    ))
    comparison = WorkbookPreservationComparator().compare_workbooks(template, path, summary)
    assert comparison["status"] == "PASS"
    assert comparison["review_required"] is False
    assert "SOURCE_FIELD_REVIEW" not in summary["status_counts"]
    assert summary["status_counts"]["EXCLUDED_RECORD"] == 1
    assert summary["normalized_record_count"] == summary["accounted_record_count"] == 2
    assert len(schedule_rows(path)) == 2
    assert {row[1] for row in schedule_rows(path)} == {"KEEP-100"}


def test_populates_current_sections_and_separate_records_preserving_template(population, tmp_path):
    template, _ = population
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    path, summary = execute(population, tmp_path, (
        "Equipment: NEW-40 Size: 400\nJOB-3 4 2026-10-05\n\f"
        "Equipment: NEW-20 Size: 200\nJOB-1 2 2026-06-08\nJOB-2 3 2026-07-13\n"
    ))
    assert summary["mode"] == "populate_template"
    assert summary["status_counts"]["CREATED_SECTION"] == 2
    assert summary["normalized_record_count"] == summary["accounted_record_count"] == 3
    assert summary["total_rows_inserted"] == 5
    assert "SECTION_NOT_FOUND" not in summary["status_counts"]
    book = openpyxl.load_workbook(path)
    sheet = book["Schedule"]
    assert [sheet.cell(row, 2).value for row in range(3, 8)] == ["NEW-20"] * 3 + ["NEW-40"] * 2
    assert sheet["F4"].value == 2
    assert sheet["G5"].value == 3
    assert sheet["I7"].value == 4
    assert sheet["E7"].value == datetime.datetime(2026, 10, 5)
    assert sheet["B3"].fill.fgColor.rgb == "00AADDCC"
    assert sheet.row_dimensions[3].height == 23
    assert sheet["F1"].value == "Overdue" and sheet["I1"].value == ">= Sep 2026"
    book.close()
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before


def test_fresh_run_rebuilds_changed_and_removed_jobs_from_current_sources(population, tmp_path):
    template, _ = population
    before = template.read_bytes()
    first_path, first_summary = execute(population, tmp_path, (
        "Equipment: KEEP-100 Size: 100\nJOB-1 4 2026-07-06\nJOB-2 7 2026-07-13\n\f"
        "Equipment: RETIRED-200 Size: 200\nJOB-3 12 2026-08-10\n"
    ), "first.xlsx")
    first_bytes = first_path.read_bytes()
    first_rows = schedule_rows(first_path)
    second_path, second_summary = execute(population, tmp_path, (
        "Equipment: KEEP-100 Size: 100\nJOB-1 9 2026-08-03\n\f"
        "Equipment: ADDED-300 Size: 300\nJOB-4 6 2026-10-05\n"
    ), "second.xlsx")
    second_rows = schedule_rows(second_path)

    assert first_summary["normalized_record_count"] == first_summary["accounted_record_count"] == 3
    assert second_summary["normalized_record_count"] == second_summary["accounted_record_count"] == 2
    assert len(first_rows) == 5
    assert len(second_rows) == 4
    assert {row[1] for row in first_rows} == {"KEEP-100", "RETIRED-200"}
    assert {row[1] for row in second_rows} == {"KEEP-100", "ADDED-300"}
    first_jobs = {row[2]: row for row in first_rows if row[2]}
    second_jobs = {row[2]: row for row in second_rows if row[2]}
    assert set(first_jobs) == {"JOB-1", "JOB-2", "JOB-3"}
    assert set(second_jobs) == {"JOB-1", "JOB-4"}
    assert first_jobs["JOB-1"][4:] == (datetime.datetime(2026, 7, 6), None, 4, None, None)
    assert second_jobs["JOB-1"][4:] == (datetime.datetime(2026, 8, 3), None, None, 9, None)
    assert second_jobs["JOB-4"][4:] == (datetime.datetime(2026, 10, 5), None, None, None, 6)
    assert template.read_bytes() == before
    assert first_path.read_bytes() == first_bytes


def test_source_reordering_and_continuation_pages_keep_the_same_workbook_records(population, tmp_path):
    template, _ = population
    before = template.read_bytes()
    first_path, first_summary = execute(population, tmp_path, (
        "Equipment: UNIT-400 Size: 400\nJOB-3 4 2026-10-05\n\f"
        "Equipment: UNIT-200 Size: 200\nJOB-1 2 2026-06-08\nJOB-2 3 2026-07-13\n"
    ), "original_order.xlsx")
    reordered_path, reordered_summary = execute(population, tmp_path, (
        "Equipment: UNIT-200 Size: 200\nJOB-2 3 2026-07-13\n\f"
        "JOB-1 2 2026-06-08\n\f"
        "Equipment: UNIT-400 Size: 400\nJOB-3 4 2026-10-05\n"
    ), "different_pages.xlsx")
    original = schedule_rows(first_path)
    reordered = schedule_rows(reordered_path)
    assert [row[1] for row in original if row[2] is None] == ["UNIT-200", "UNIT-400"]
    assert [row[1] for row in reordered if row[2] is None] == ["UNIT-200", "UNIT-400"]
    assert sorted(original, key=lambda row: (row[1], row[2] or "")) == sorted(
        reordered, key=lambda row: (row[1], row[2] or ""),
    )
    for summary in (first_summary, reordered_summary):
        assert summary["normalized_record_count"] == summary["accounted_record_count"] == 3
        assert summary["total_rows_inserted"] == 5
    assert template.read_bytes() == before


def test_configured_dependencies_appear_before_parent_even_when_size_sorts_later(dependency_population, tmp_path):
    path, summary = execute(dependency_population, tmp_path, (
        "Equipment: ASSEMBLY-100 Size: 100 Dependency: COMPONENT-900\nJOB-1 4 2026-07-06\n\f"
        "Equipment: INDEPENDENT-200 Size: 200\nJOB-2 7 2026-08-03\n\f"
        "Equipment: COMPONENT-900 Size: 900\nJOB-3 2 2026-07-13\n"
    ))
    rows = schedule_rows(path)
    assert [row[1] for row in rows if row[2] is None] == [
        "COMPONENT-900", "ASSEMBLY-100", "INDEPENDENT-200",
    ]
    assert summary["total_rows_inserted"] == 6
    assert summary["normalized_record_count"] == summary["accounted_record_count"] == 3


def test_dependency_order_does_not_reintroduce_excluded_sections(dependency_population, tmp_path):
    _, profile = dependency_population
    parameters = profile["parameters"]
    parameters["section_selection"] = {"field": "id", "operator": "equals", "value": "ASSEMBLY-100"}
    parameters["section_order"]["dependencies_before"]["missing"] = "ignore"
    path, summary = execute(dependency_population, tmp_path, (
        "Equipment: ASSEMBLY-100 Size: 100 Dependency: EXCLUDED-900\nJOB-1 4 2026-07-06\n\f"
        "Equipment: EXCLUDED-900 Size: 900\nJOB-2 7 2026-08-03\nJOB-3 2 2026-07-13\n"
    ))
    rows = schedule_rows(path)
    assert {row[1] for row in rows} == {"ASSEMBLY-100"}
    assert summary["total_rows_inserted"] == 2
    assert summary["status_counts"]["EXCLUDED_SECTION"] == 1
    assert summary["status_counts"]["EXCLUDED_RECORD"] == 2
    assert summary["normalized_record_count"] == summary["accounted_record_count"] == 3


def test_dependency_order_refuses_duplicate_selected_section_keys(dependency_population, tmp_path):
    template, _ = dependency_population
    before = template.read_bytes()
    with pytest.raises(ValueError, match="unique selected section keys"):
        execute(dependency_population, tmp_path, (
            "Equipment: DUPLICATE-100 Size: 100\nJOB-1 4 2026-07-06\n\f"
            "Equipment: DUPLICATE-100 Size: 100\nJOB-2 7 2026-08-03\n"
        ))
    assert not (tmp_path / "result.xlsx").exists()
    assert template.read_bytes() == before


def test_dependency_order_refuses_cycles_without_publishing_workbook(dependency_population, tmp_path):
    template, _ = dependency_population
    before = template.read_bytes()
    with pytest.raises(ValueError, match="dependency cycle"):
        execute(dependency_population, tmp_path, (
            "Equipment: ASSEMBLY-100 Size: 100 Dependency: COMPONENT-200\nJOB-1 4 2026-07-06\n\f"
            "Equipment: COMPONENT-200 Size: 200 Dependency: ASSEMBLY-100\nJOB-2 7 2026-08-03\n"
        ))
    assert not (tmp_path / "result.xlsx").exists()
    assert template.read_bytes() == before


def test_new_materials_and_new_months_do_not_require_profile_changes(population, tmp_path):
    template, _ = population
    book = openpyxl.load_workbook(template)
    sheet = book["Schedule"]
    sheet["G1"] = datetime.date(2028, 1, 1)
    sheet["H1"] = datetime.date(2028, 2, 1)
    sheet["I1"] = ">= Mar 2028"
    book.save(template)
    book.close()
    path, summary = execute(population, tmp_path, (
        "Equipment: UNSEEN-900 Size: 900\nJOB-99 17 2028-02-07\n\f"
        "Equipment: UNSEEN-100 Size: 100\nJOB-98 8 2028-04-03\n"
    ))
    assert summary["total_rows_inserted"] == 4
    book = openpyxl.load_workbook(path)
    sheet = book["Schedule"]
    assert sheet["B3"].value == "UNSEEN-100" and sheet["I4"].value == 8
    assert sheet["B5"].value == "UNSEEN-900" and sheet["H6"].value == 17
    book.close()


def test_population_refuses_populated_input_instead_of_duplicating_or_overwriting(population, tmp_path):
    template, _ = population
    book = openpyxl.load_workbook(template)
    book["Schedule"]["A8"] = "Planner note that must survive"
    book.save(template)
    book.close()
    before = template.read_bytes()
    with pytest.raises(ValueError, match="requires an empty data area"):
        execute(population, tmp_path, "Equipment: NEW-1 Size: 100\nJOB-1 4 2026-07-06\n")
    assert not (tmp_path / "result.xlsx").exists()
    assert template.read_bytes() == before


def test_template_with_only_headers_and_no_formatted_body(population, tmp_path):
    template, _ = population
    book = openpyxl.load_workbook(template)
    book["Schedule"].delete_rows(2, book["Schedule"].max_row)
    book.save(template)
    book.close()
    path, summary = execute(population, tmp_path, "Equipment: NEW-1 Size: 100\nJOB-1 4 2026-07-06\n")
    assert summary["total_rows_inserted"] == 2
    book = openpyxl.load_workbook(path)
    assert book["Schedule"]["B3"].value == "NEW-1"
    assert book["Schedule"]["G4"].value == 4
    book.close()


def test_default_update_mode_still_requires_existing_sections(population, tmp_path):
    _, profile = population
    del profile["parameters"]["workbook_update"]["mode"]
    _, summary = execute(population, tmp_path, "Equipment: NEW-1 Size: 100\nJOB-1 4 2026-07-06\n")
    assert summary["total_rows_inserted"] == 0
    assert summary["status_counts"]["SECTION_NOT_FOUND"] == 2


def test_unknown_planning_header_still_fails(population, tmp_path):
    template, _ = population
    book = openpyxl.load_workbook(template)
    book["Schedule"]["J1"] = "Unmapped measure"
    book.save(template)
    book.close()
    with pytest.raises(ValueError, match="Unsupported planning header"):
        execute(population, tmp_path, "Equipment: NEW-1 Size: 100\nJOB-1 4 2026-07-06\n")


def test_equivalent_style_table_entries_and_blank_values_round_trip(population, tmp_path):
    template, profile = population
    # Real Excel files may contain duplicate alignment entries with different IDs.
    # A save deduplicates these; that must not count as changing cell formatting.
    with ZipFile(template) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    styles = ElementTree.fromstring(files["xl/styles.xml"])
    alignments = styles.find("s:cellXfs", ns)
    original = alignments[1]
    original.append(ElementTree.Element("{" + ns["s"] + "}alignment", horizontal="left"))
    duplicate = copy.deepcopy(original)
    alignments.append(duplicate)
    alignments.set("count", str(len(alignments)))
    files["xl/styles.xml"] = ElementTree.tostring(styles)
    sheet = ElementTree.fromstring(files["xl/worksheets/sheet1.xml"])
    sheet.find(".//s:c[@r='B3']", ns).set("s", str(len(alignments) - 1))
    files["xl/worksheets/sheet1.xml"] = ElementTree.tostring(sheet)
    with ZipFile(template, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    profile["parameters"]["workbook_update"]["template_section_values"]["4"] = {"literal": ""}
    path, summary = execute(population, tmp_path, "Equipment: FRESH-1 Size: 100\nJOB-1 4 2026-07-06\n")
    assert summary["total_rows_inserted"] == 2
    book = openpyxl.load_workbook(path)
    assert book["Schedule"]["B3"].alignment.horizontal == "left"
    assert book["Schedule"]["D3"].value is None
    book.close()


@pytest.mark.parametrize("change", [
    {"mode": []}, {"backlog_headers": "Overdue"}, {"future_header_patterns": [False]},
    {"template_section_values": {"0": {"literal": "bad"}}}, {"template_section_sort": [{}] * 9},
])
def test_population_configuration_is_validated(population, change):
    _, profile = population
    modified = copy.deepcopy(profile)
    modified["parameters"]["workbook_update"].update(change)
    with pytest.raises(ValueError):
        validate_workbook_update_profile(modified)
